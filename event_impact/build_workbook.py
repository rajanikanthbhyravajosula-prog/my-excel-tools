"""Build Event_Impact_Analyzer.xlsx from the CSVs produced by analyze.py.

Raw measurements (per event, per stock) are written as values; every summary,
ranking, score and the target adjuster are live Excel formulas that refer to
them, so changing a setting (yellow cell) recalculates the whole workbook.
"""
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
DATA = HERE / "data"
OUT = HERE / "Event_Impact_Analyzer.xlsx"
FII_SOURCE = HERE / "fii_dii_daily.csv"  # exported from the FII_DII_Daily sheet of the shares worksheet
FII_ROWS = 1000  # pre-filled formula rows on the FII_DII sheet

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F3864")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_SUB = PatternFill("solid", fgColor="D9E1F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_OUT = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = '0.0%;[Red]-0.0%;"-"'
NUM1 = '0.0;[Red]-0.0;"-"'
NUM0 = '#,##0;[Red]-#,##0;"-"'
RS = '"₹"#,##0.00'
CR = '#,##0;[Red](#,##0);"-"'

RED_FILL = PatternFill("solid", fgColor="F8CBAD")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREY_FILL = PatternFill("solid", fgColor="EDEDED")


def scale3():
    return ColorScaleRule(start_type="min", start_color="F8696B", mid_type="num", mid_value=0,
                          mid_color="FFFFFF", end_type="max", end_color="63BE7B")


def header(ws, row, labels, col=1, fill=FILL_HEAD, font=F_HEAD, height=30):
    for i, text in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=text)
        c.font, c.fill, c.border = font, fill, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = height


def title(ws, text, note=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if note:
        ws["A2"] = note
        ws["A2"].font = F_NOTE


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def style_range(ws, rng, fmt=None, font=F_BASE, border=True):
    for row in ws[rng]:
        for c in row:
            c.font = font
            if fmt:
                c.number_format = fmt
            if border:
                c.border = BOX


def mood_rules(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fear"'], fill=RED_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Greed"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Calm"'], fill=GREY_FILL))


def nz(v):
    return None if pd.isna(v) else v


# ---------------------------------------------------------------- settings
SETTINGS = [  # (row, label, value, format, explanation)
    (4, "Fear threshold: Nifty worst fall within 20 days of the event", -0.05, PCT,
     "An event is tagged FEAR when Nifty fell at least this much (and the fall was bigger than the rise)."),
    (5, "Greed threshold: Nifty best rise within 20 days of the event", 0.05, PCT,
     "An event is tagged GREED when Nifty rose at least this much. Everything else is CALM."),
    (6, "Points scale: average move worth 1 point", 0.005, PCT,
     "0.5% means an average 20-day move of +3% scores +6 points (capped at +/-10)."),
    (7, "Minimum events before trusting a stock's own history", 5, "0",
     "Below this, the Target Adjuster uses Event Beta x Nifty instead of the stock's own average."),
    (8, "Short-term target weight", 1.0, "0.00",
     "1.00 = apply the full expected event move to your short-term target."),
    (9, "Long-term target weight", 0.25, "0.00",
     "Long-term targets should follow fundamentals; 0.25 applies only a quarter of the event move."),
    (10, "Stop-loss buffer multiplier", 1.0, "0.00",
     "1.00 = stop-loss at the typical worst dip after this kind of event; 1.2 = 20% wider."),
    (11, "FII/DII 'big flow' threshold (Rs crore)", 3000, "#,##0",
     "Used on the FII_DII sheet to split big buying days from big selling days."),
    (12, "Minimum crises before ranking a stock", 3, "0",
     "Crisis_Scorecard ranks only stocks measured in at least this many crisis periods."),
]
S = {k: f"Settings!$B${r}" for k, r in zip(
    ["fear", "greed", "scale", "minev", "stw", "ltw", "slm", "fii", "mincr"], [r for r, *_ in SETTINGS])}


def build_settings(wb):
    ws = wb.create_sheet("Settings")
    title(ws, "Settings — change the yellow cells; the whole workbook recalculates")
    header(ws, 3, ["Setting", "Value", "What it means"])
    for r, label, val, fmt, expl in SETTINGS:
        ws.cell(r, 1, label).font = F_BASE
        c = ws.cell(r, 2, val)
        c.font, c.fill, c.number_format, c.border = F_INPUT, FILL_INPUT, fmt, BOX
        ws.cell(r, 3, expl).font = F_NOTE
    widths(ws, {"A": 58, "B": 12, "C": 100})


# ---------------------------------------------------------------- calendar + nifty impact
NI_COLS = [  # (csv column, header, format)
    ("Event_ID", "Event ID", None), ("Trading_Day", "Trading day (T0)", "dd-mmm-yyyy"),
    ("Category", "Category", None), ("News_Tone", "News tone", None), ("Surprise", "Surprise?", None),
    ("Headline", "Headline", None), ("Index", "Index used", None),
    ("Pre20", "BEFORE: 20-day run-up", PCT), ("Pre5", "BEFORE: 5-day run-up", PCT),
    ("Day0", "DURING: event-day move", PCT), ("Day0_Low", "DURING: intraday worst", PCT),
    ("Post5", "AFTER: next 5 days", PCT), ("Post20", "AFTER: next 20 days", PCT),
    ("Post60", "AFTER: next 60 days", PCT), ("Total20", "Total move T-1 to T+20", PCT),
    ("MaxFall20", "Worst fall within 20 days", PCT), ("MaxRise20", "Best rise within 20 days", PCT),
    ("Recovery_Days", "Days to recover pre-event level", NUM0),
    ("VIX_Before", "India VIX before", NUM1), ("VIX_Peak20", "VIX peak (20 days)", NUM1),
    ("VIX_Change", "VIX jump", PCT), ("VIX_Settle_Days", "Days for VIX to settle", NUM0),
    ("Ret20_Before", "Buy 5 days BEFORE: 20-day return", PCT),
    ("Reward_Before", "Buy BEFORE: best gain", PCT), ("Risk_Before", "Buy BEFORE: worst dip", PCT),
    ("Ret20_EventDay", "Buy ON EVENT DAY: 20-day return", PCT),
    ("Reward_EventDay", "Buy ON EVENT DAY: best gain", PCT), ("Risk_EventDay", "Buy ON EVENT DAY: worst dip", PCT),
    ("Ret20_After", "Buy 5 days AFTER: 20-day return", PCT),
    ("Reward_After", "Buy AFTER: best gain", PCT), ("Risk_After", "Buy AFTER: worst dip", PCT),
]
NI = {name: L(i + 1) for i, (name, _, _) in enumerate(NI_COLS)}
NI["Mood"] = L(len(NI_COLS) + 1)
NI["Up20"] = L(len(NI_COLS) + 2)
NI["PanicBefore"] = L(len(NI_COLS) + 3)
NI["PanicPeak"] = L(len(NI_COLS) + 4)
MDC = {}  # Market_Daily column letters, filled by build_market_daily


def build_calendar(wb, ni):
    ws = wb.create_sheet("Event_Calendar")
    title(ws, "Event Calendar — the news behind each event",
          "Dates and news summaries were compiled for this analysis; please verify them. Crude Spike/Crash rows "
          "are detected automatically from Brent prices. To add an event, add a row to events.csv and re-run the scripts.")
    cols = ["Event ID", "Event date", "Trading day (T0)", "Category", "News tone", "Surprise?", "Headline",
            "What happened / what the market expected"]
    header(ws, 4, cols)
    for r, row in enumerate(ni.itertuples(index=False), start=5):
        vals = [row.Event_ID, pd.Timestamp(row.Date).to_pydatetime(), pd.Timestamp(row.Trading_Day).to_pydatetime(),
                row.Category, row.News_Tone, row.Surprise, row.Headline, row.Context]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = F_BASE, BOX
            if c in (2, 3):
                cell.number_format = "dd-mmm-yyyy"
            if c == 8:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:H{4 + len(ni)}"
    widths(ws, {"A": 9, "B": 12, "C": 12, "D": 17, "E": 10, "F": 9, "G": 42, "H": 90})


def build_nifty_impact(wb, ni):
    ws = wb.create_sheet("Nifty_Impact")
    heads = [h for _, h, _ in NI_COLS] + ["Market mood (formula)", "Higher after 20 days? (1=yes)",
                                         "Panic score day before (0-100)", "Panic peak within 20 days"]
    header(ws, 1, heads, height=45)
    n = len(ni)
    for r, row in enumerate(ni.to_dict("records"), start=2):
        for c, (name, _, fmt) in enumerate(NI_COLS, start=1):
            v = row[name]
            if name == "Trading_Day":
                v = pd.Timestamp(v).to_pydatetime()
            cell = ws.cell(r, c, nz(v))
            cell.font, cell.border = F_BASE, BOX
            if fmt:
                cell.number_format = fmt
        f, g = NI["MaxFall20"], NI["MaxRise20"]
        ws[f"{NI['Mood']}{r}"] = (f'=IF({f}{r}="","",IF(AND({f}{r}<={S["fear"]},-{f}{r}>={g}{r}),"Fear",'
                                  f'IF({g}{r}>={S["greed"]},"Greed","Calm")))')
        t = NI["Total20"]
        ws[f"{NI['Up20']}{r}"] = f'=IF({t}{r}="","",IF({t}{r}>0,1,0))'
        pc, dc = f"Market_Daily!${MDC['Panic']}:${MDC['Panic']}", "Market_Daily!$A:$A"
        m = f"MATCH(B{r},{dc},0)"
        ws[f"{NI['PanicBefore']}{r}"] = f'=IFERROR(IF({m}<3,"",INDEX({pc},{m}-1)),"")'
        ws[f"{NI['PanicPeak']}{r}"] = f'=IFERROR(MAX(INDEX({pc},{m}):INDEX({pc},{m}+20)),"")'
        for k in ("Mood", "Up20", "PanicBefore", "PanicPeak"):
            ws[f"{NI[k]}{r}"].font, ws[f"{NI[k]}{r}"].border = F_BASE, BOX
        for k in ("PanicBefore", "PanicPeak"):
            ws[f"{NI[k]}{r}"].number_format = "0"
    mood_rules(ws, f"{NI['Mood']}2:{NI['Mood']}{n + 1}")
    for k in ("Day0", "Total20", "Post60"):
        ws.conditional_formatting.add(f"{NI[k]}2:{NI[k]}{n + 1}", scale3())
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{NI['PanicPeak']}{n + 1}"
    ws.conditional_formatting.add(f"{NI['PanicBefore']}2:{NI['PanicPeak']}{n + 1}", panic_scale())
    for i in range(1, len(heads) + 1):
        ws.column_dimensions[L(i)].width = 11
    widths(ws, {"C": 16, "F": 40, NI["Mood"]: 12})
    return n + 1  # last data row


def ni_rng(key, last):
    return f"Nifty_Impact!${NI[key]}$2:${NI[key]}${last}"


# ---------------------------------------------------------------- category summary
def build_category_summary(wb, categories, last):
    ws = wb.create_sheet("Category_Summary")
    title(ws, "Category Summary — how Nifty behaved around each type of event (all formulas)",
          "Averages over every event in the category. Phases: BEFORE = run-up into the event, DURING = event day, "
          "AFTER = following days. R:R = average best gain / average worst dip in the 20 days after buying.")
    heads = ["Category", "Events", "BEFORE: avg 20-day run-up", "DURING: avg event-day move",
             "AFTER: avg next 5 days", "Avg total move (T-1 to T+20)", "AFTER: avg next 60 days",
             "% of events Nifty higher after 20 days", "Avg worst fall (20 days)", "Avg best rise (20 days)",
             "Avg days to recover", "Avg VIX jump", "Avg days for VIX to settle", "Fear events", "Greed events",
             "Buy BEFORE: avg 20-day return", "Buy BEFORE: R:R", "Buy ON DAY: avg 20-day return", "Buy ON DAY: R:R",
             "Buy AFTER: avg 20-day return", "Buy AFTER: R:R", "Best entry (history)",
             "Direction points (-10 to +10)", "Risk points (0 to 10)", "Confidence", "Signal"]
    header(ws, 4, heads, height=60)
    cat = ni_rng("Category", last)

    def avg(key, r):
        return f'=IFERROR(AVERAGEIFS({ni_rng(key, last)},{cat},$A{r}),"")'

    def rr(label, r):
        return (f'=IFERROR(AVERAGEIFS({ni_rng("Reward_" + label, last)},{cat},$A{r})/'
                f'AVERAGEIFS({ni_rng("Risk_" + label, last)},{cat},$A{r}),"")')

    for r, name in enumerate(categories, start=5):
        ws.cell(r, 1, name)
        ws.cell(r, 2, f"=COUNTIF({cat},$A{r})")
        for c, key in zip(range(3, 14), ["Pre20", "Day0", "Post5", "Total20", "Post60", "Up20", "MaxFall20",
                                         "MaxRise20", "Recovery_Days", "VIX_Change", "VIX_Settle_Days"]):
            ws.cell(r, c, avg(key, r))
        ws.cell(r, 14, f'=COUNTIFS({cat},$A{r},{ni_rng("Mood", last)},"Fear")')
        ws.cell(r, 15, f'=COUNTIFS({cat},$A{r},{ni_rng("Mood", last)},"Greed")')
        ws.cell(r, 16, avg("Ret20_Before", r))
        ws.cell(r, 17, rr("Before", r))
        ws.cell(r, 18, avg("Ret20_EventDay", r))
        ws.cell(r, 19, rr("EventDay", r))
        ws.cell(r, 20, avg("Ret20_After", r))
        ws.cell(r, 21, rr("After", r))
        ws.cell(r, 22, f'=IF(P{r}="","",IF(AND(P{r}>=R{r},P{r}>=T{r}),"Buy before",'
                       f'IF(R{r}>=T{r},"Buy on event day","Buy 5 days after")))')
        ws.cell(r, 23, f'=IF(F{r}="","",MAX(-10,MIN(10,ROUND(F{r}/{S["scale"]},0))))')
        ws.cell(r, 24, f'=IF(I{r}="","",MAX(0,MIN(10,ROUND(-I{r}/(2*{S["scale"]}),0))))')
        ws.cell(r, 25, f'=IF(B{r}>=10,"High",IF(B{r}>=5,"Medium","Low"))')
        ws.cell(r, 26, f'=IF(W{r}="","",IF(W{r}>=2,"Bullish bias",IF(W{r}<=-2,"Bearish bias","Neutral")))')
    end = 4 + len(categories)
    style_range(ws, f"A5:Z{end}")
    for col in "CDEFGHIJLPRT":
        style_range(ws, f"{col}5:{col}{end}", PCT)
    for col in "KM":
        style_range(ws, f"{col}5:{col}{end}", NUM0)
    for col in "QSU":
        style_range(ws, f"{col}5:{col}{end}", "0.00")
    for col in "DFG":
        ws.conditional_formatting.add(f"{col}5:{col}{end}", scale3())
    ws.conditional_formatting.add(f"W5:W{end}", scale3())
    for r in range(5, end + 1):
        ws.cell(r, 1).font = F_BOLD

    # ---- news tone x surprise
    r0 = end + 3
    ws.cell(r0 - 1, 1, "How the market reacts to POSITIVE vs NEGATIVE news (and whether it was a surprise)").font = F_BOLD
    header(ws, r0, ["News tone", "Surprise?", "Events", "Avg event-day move", "Avg total move (20 days)",
                    "% higher after 20 days", "Avg worst fall", "Avg days to recover"], height=45)
    tone, sur = ni_rng("News_Tone", last), ni_rng("Surprise", last)
    r = r0 + 1
    for t in ["Positive", "Negative", "Mixed", "Neutral"]:
        for s in ["Yes", "No", "Partly"]:
            ws.cell(r, 1, t)
            ws.cell(r, 2, s)
            ws.cell(r, 3, f"=COUNTIFS({tone},$A{r},{sur},$B{r})")
            for c, key in zip(range(4, 9), ["Day0", "Total20", "Up20", "MaxFall20", "Recovery_Days"]):
                ws.cell(r, c, f'=IFERROR(AVERAGEIFS({ni_rng(key, last)},{tone},$A{r},{sur},$B{r}),"")')
            r += 1
    style_range(ws, f"A{r0 + 1}:H{r - 1}")
    for col in "DEFG":
        style_range(ws, f"{col}{r0 + 1}:{col}{r - 1}", PCT)
    style_range(ws, f"H{r0 + 1}:H{r - 1}", NUM0)
    ws.conditional_formatting.add(f"E{r0 + 1}:E{r - 1}", scale3())

    # ---- mood summary: who made the money
    m0 = r + 2
    ws.cell(m0 - 1, 1, "FEAR vs GREED events — who made the money? (buyer's average 20-day result by entry timing)").font = F_BOLD
    header(ws, m0, ["Market mood", "Events", "Avg worst fall", "Avg days to recover", "Avg days for VIX to settle",
                    "Buy BEFORE: 20-day return", "Buy ON DAY: 20-day return", "Buy AFTER: 20-day return",
                    "Buy ON DAY: R:R", "Who made the money"], height=45)
    mood = ni_rng("Mood", last)
    for i, m in enumerate(["Fear", "Greed", "Calm"]):
        rr_ = m0 + 1 + i
        ws.cell(rr_, 1, m)
        ws.cell(rr_, 2, f'=COUNTIF({mood},$A{rr_})')
        for c, key in zip(range(3, 9), ["MaxFall20", "Recovery_Days", "VIX_Settle_Days", "Ret20_Before",
                                        "Ret20_EventDay", "Ret20_After"]):
            ws.cell(rr_, c, f'=IFERROR(AVERAGEIFS({ni_rng(key, last)},{mood},$A{rr_}),"")')
        ws.cell(rr_, 9, f'=IFERROR(AVERAGEIFS({ni_rng("Reward_EventDay", last)},{mood},$A{rr_})/'
                        f'AVERAGEIFS({ni_rng("Risk_EventDay", last)},{mood},$A{rr_}),"")')
        ws.cell(rr_, 10, f'=IF(G{rr_}="","",IF(MAX(F{rr_}:H{rr_})<0,"Sellers before the event / cash holders",'
                         f'IF(AND(H{rr_}>=G{rr_},H{rr_}>=F{rr_}),"Patient buyers who waited 5 days",'
                         f'IF(G{rr_}>=F{rr_},"Buyers on the event day","Early buyers (positioned before)"))))')
    style_range(ws, f"A{m0 + 1}:J{m0 + 3}")
    for col in "CFGH":
        style_range(ws, f"{col}{m0 + 1}:{col}{m0 + 3}", PCT)
    for col in "DE":
        style_range(ws, f"{col}{m0 + 1}:{col}{m0 + 3}", NUM0)
    style_range(ws, f"I{m0 + 1}:I{m0 + 3}", "0.00")
    mood_rules(ws, f"A{m0 + 1}:A{m0 + 3}")

    ws.freeze_panes = "B5"
    widths(ws, {"A": 18, "V": 17, "Y": 11, "Z": 13, "J": 13})
    for i in range(2, 27):
        if L(i) not in ("V", "Y", "Z", "J"):
            ws.column_dimensions[L(i)].width = 11
    return f"Category_Summary!$A$5:$A${end}"


# ---------------------------------------------------------------- stock level
SE_COLS = ["Stock", "Event ID", "Trading day", "Category", "Industry", "Cap", "Market mood",
           "Stock event-day move", "Nifty event-day move", "Excess vs Nifty (day)",
           "Stock total move (20 days)", "Nifty total move (20 days)", "Excess vs Nifty (20 days)",
           "Stock worst fall (20 days)", "Stock best rise (20 days)", "Days to recover",
           "Stock next 60 days", "Buy on day: best gain", "Buy on day: worst dip", "Buy on day: 20-day return"]


def build_stock_events(wb, se, last_ni):
    ws = wb.create_sheet("Stock_Event_Data")
    header(ws, 1, SE_COLS, height=45)
    for r, row in enumerate(se.itertuples(index=False), start=2):
        vals = {1: row.Stock, 2: row.Event_ID, 3: pd.Timestamp(row.Trading_Day).to_pydatetime(),
                4: row.Category, 5: row.Industry, 6: row.Cap,
                8: nz(row.Day0), 9: nz(row.Nifty_Day0), 11: nz(row.Total20), 12: nz(row.Nifty_Total20),
                14: nz(row.MaxFall20), 15: nz(row.MaxRise20), 16: nz(row.Recovery_Days), 17: nz(row.Post60),
                18: nz(row.Reward_EventDay), 19: nz(row.Risk_EventDay), 20: nz(row.Ret20_EventDay)}
        for c, v in vals.items():
            ws.cell(r, c, v)
        ws.cell(r, 7, f'=INDEX({ni_rng("Mood", last_ni)},MATCH(B{r},{ni_rng("Event_ID", last_ni)},0))')
        ws.cell(r, 10, f'=IF(OR(H{r}="",I{r}=""),"",H{r}-I{r})')
        ws.cell(r, 13, f'=IF(OR(K{r}="",L{r}=""),"",K{r}-L{r})')
    last = len(se) + 1
    for col, fmt in (("C", "dd-mmm-yyyy"), ("P", NUM0)):
        for c in ws[f"{col}2:{col}{last}"]:
            c[0].number_format = fmt
    for col in "HIJKLMNOQRST":
        for c in ws[f"{col}2:{col}{last}"]:
            c[0].number_format = PCT
    mood_rules(ws, f"G2:G{last}")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:T{last}"
    for i in range(1, 21):
        ws.column_dimensions[L(i)].width = 11
    widths(ws, {"D": 16, "E": 18})
    return last


def se_rng(col, last):
    return f"Stock_Event_Data!${col}$2:${col}${last}"


def build_scorecard(wb, watch, last_se):
    ws = wb.create_sheet("Stock_Scorecard")
    title(ws, "Stock Scorecard — how each watchlist stock behaves around events (all formulas)",
          "FEAR columns use only events where Nifty fell (see Settings). Event Beta = how much the stock moves for "
          "every 1% Nifty move around events. Lower 'score' = better rank.")
    heads = ["Stock", "Name", "Industry", "Cap", "Events measured", "First data row", "Fear events",
             "Avg event-day excess vs Nifty", "FEAR: avg 20-day move", "FEAR: avg excess vs Nifty",
             "FEAR: avg worst fall", "FEAR: avg days to recover", "GREED: avg 20-day move",
             "GREED: avg excess vs Nifty", "Event Beta", "Event R:R (buy on day)",
             "Resilience score", "Resilience rank (1 = best)", "Impact score", "Impact rank (1 = most hit)"]
    header(ws, 4, heads, height=60)
    stk, mood = se_rng("A", last_se), se_rng("G", last_se)
    n = len(watch)
    end = 4 + n
    for r, w in enumerate(watch.itertuples(index=False), start=5):
        ws.cell(r, 1, w.NSE_Code).font = F_BOLD
        ws.cell(r, 2, w.Name)
        ws.cell(r, 3, w.Industry)
        ws.cell(r, 4, w.Cap)
        ws.cell(r, 5, f"=COUNTIF({stk},$A{r})")
        ws.cell(r, 6, f'=IFERROR(MATCH($A{r},{stk},0)+1,"")')
        ws.cell(r, 7, f'=COUNTIFS({stk},$A{r},{mood},"Fear")')
        ws.cell(r, 8, f'=IFERROR(AVERAGEIFS({se_rng("J", last_se)},{stk},$A{r}),"")')
        for c, col in ((9, "K"), (10, "M"), (11, "N"), (12, "P")):
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({se_rng(col, last_se)},{stk},$A{r},{mood},"Fear"),"")')
        for c, col in ((13, "K"), (14, "M")):
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({se_rng(col, last_se)},{stk},$A{r},{mood},"Greed"),"")')
        ws.cell(r, 15, f'=IF(E{r}<3,"",IFERROR(SLOPE(INDEX(Stock_Event_Data!$K:$K,F{r}):INDEX(Stock_Event_Data!$K:$K,F{r}+E{r}-1),'
                       f'INDEX(Stock_Event_Data!$L:$L,F{r}):INDEX(Stock_Event_Data!$L:$L,F{r}+E{r}-1)),""))')
        ws.cell(r, 16, f'=IFERROR(AVERAGEIFS({se_rng("R", last_se)},{stk},$A{r})/AVERAGEIFS({se_rng("S", last_se)},{stk},$A{r}),"")')
        ok = f'OR(G{r}<{S["minev"]},J{r}="",K{r}="",L{r}="",O{r}="")'
        ws.cell(r, 17, f'=IF({ok},"",RANK(J{r},$J$5:$J${end},0)+RANK(K{r},$K$5:$K${end},0)'
                       f'+RANK(L{r},$L$5:$L${end},1)+ROW()/100000)')
        ws.cell(r, 18, f'=IF(Q{r}="","",RANK(Q{r},$Q$5:$Q${end},1))')
        ws.cell(r, 19, f'=IF({ok},"",RANK(K{r},$K$5:$K${end},1)+RANK(O{r},$O$5:$O${end},0)+ROW()/100000)')
        ws.cell(r, 20, f'=IF(S{r}="","",RANK(S{r},$S$5:$S${end},1))')
    style_range(ws, f"A5:T{end}")
    for r in range(5, end + 1):
        ws.cell(r, 1).font = F_BOLD
    for col in "HIJKMN":
        style_range(ws, f"{col}5:{col}{end}", PCT)
    style_range(ws, f"L5:L{end}", NUM0)
    for col in "OP":
        style_range(ws, f"{col}5:{col}{end}", "0.00")
    for col in "QS":
        style_range(ws, f"{col}5:{col}{end}", "0.0")
    for col in "JKN":
        ws.conditional_formatting.add(f"{col}5:{col}{end}", scale3())
    ws.freeze_panes = "B5"
    ws.auto_filter.ref = f"A4:T{end}"
    for i in range(1, 21):
        ws.column_dimensions[L(i)].width = 11
    widths(ws, {"B": 20, "C": 18, "F": 7})
    ws.column_dimensions["F"].hidden = True
    return end


def build_top_lists(wb, end_sc, top=25):
    ws = wb.create_sheet("Top_Lists")
    title(ws, "Top Lists — most resilient vs most impacted stocks in FEAR events (formulas)",
          "Resilient = fell less than Nifty, smaller dips, quicker recovery. Most impacted = deepest dips and highest Event Beta. "
          "Only stocks with enough FEAR events (Settings) are ranked.")
    sc = lambda c: f"Stock_Scorecard!${c}$5:${c}${end_sc}"  # noqa: E731
    blocks = [("MOST RESILIENT — performed best even in fear", "Q", 4),
              ("MOST IMPACTED — hit hardest in fear", "S", 8 + top)]
    for label, score_col, r0 in blocks:
        ws.cell(r0 - 1, 1, label).font = F_BOLD
        header(ws, r0, ["Rank", "Stock", "Name", "Industry", "Cap", "Fear events", "FEAR: avg 20-day move",
                        "FEAR: excess vs Nifty", "FEAR: avg worst fall", "FEAR: days to recover",
                        "GREED: avg 20-day move", "Event Beta"], height=45)
        for k in range(1, top + 1):
            r = r0 + k
            ws.cell(r, 1, k)
            match = f"MATCH(SMALL({sc(score_col)},$A{r}),{sc(score_col)},0)"
            for c, col in zip(range(2, 13), "ABCDGIJKLMO"):
                ws.cell(r, c, f'=IFERROR(INDEX({sc(col)},{match}),"")')
        style_range(ws, f"A{r0 + 1}:L{r0 + top}")
        for col in "GHIK":
            style_range(ws, f"{col}{r0 + 1}:{col}{r0 + top}", PCT)
        style_range(ws, f"J{r0 + 1}:J{r0 + top}", NUM0)
        style_range(ws, f"L{r0 + 1}:L{r0 + top}", "0.00")
    widths(ws, {"A": 6, "B": 13, "C": 20, "D": 18, "E": 10})
    for col in "FGHIJKL":
        ws.column_dimensions[col].width = 12


def build_sector(wb, industries, categories, last_se):
    ws = wb.create_sheet("Sector_Impact")
    title(ws, "Sector Impact — average 20-day move of watchlist stocks in each industry, by event type (formulas)",
          "Green = rose on average after this type of event, red = fell. Built from Stock_Event_Data.")
    heads = ["Industry", "Stock-events"] + categories + ["FEAR: avg 20-day move", "FEAR: avg worst fall",
                                                        "FEAR: days to recover", "GREED: avg 20-day move"]
    header(ws, 4, heads, height=60)
    ind, cat, mood = se_rng("E", last_se), se_rng("D", last_se), se_rng("G", last_se)
    nc = len(categories)
    for r, name in enumerate(industries, start=5):
        ws.cell(r, 1, name).font = F_BOLD
        ws.cell(r, 2, f"=COUNTIF({ind},$A{r})")
        for j in range(nc):
            col = L(3 + j)
            ws.cell(r, 3 + j, f'=IFERROR(AVERAGEIFS({se_rng("K", last_se)},{ind},$A{r},{cat},{col}$4),"")')
        base = 3 + nc
        ws.cell(r, base, f'=IFERROR(AVERAGEIFS({se_rng("K", last_se)},{ind},$A{r},{mood},"Fear"),"")')
        ws.cell(r, base + 1, f'=IFERROR(AVERAGEIFS({se_rng("N", last_se)},{ind},$A{r},{mood},"Fear"),"")')
        ws.cell(r, base + 2, f'=IFERROR(AVERAGEIFS({se_rng("P", last_se)},{ind},$A{r},{mood},"Fear"),"")')
        ws.cell(r, base + 3, f'=IFERROR(AVERAGEIFS({se_rng("K", last_se)},{ind},$A{r},{mood},"Greed"),"")')
    end = 4 + len(industries)
    style_range(ws, f"A5:{L(base + 3)}{end}", PCT)
    style_range(ws, f"B5:B{end}", NUM0)
    style_range(ws, f"{L(base + 2)}5:{L(base + 2)}{end}", NUM0)
    for r in range(5, end + 1):
        ws.cell(r, 1).font = F_BOLD
    ws.conditional_formatting.add(f"C5:{L(base - 1)}{end}", scale3())
    for k in (0, 1, 3):
        c = L(base + k)
        ws.conditional_formatting.add(f"{c}5:{c}{end}", scale3())
    ws.freeze_panes = "C5"
    widths(ws, {"A": 22, "B": 9})
    for i in range(3, base + 4):
        ws.column_dimensions[L(i)].width = 11


# ---------------------------------------------------------------- crude + daily data
PANIC_COMPONENTS = [  # (Market_Daily column, label, calm value -> score 0, fear value -> score 100, meaning)
    ("VIX", "India VIX level", 12, 40, "Expected volatility. Below 12 = complacent; 40+ = panic (2008: 85, 2020: 84)."),
    ("Drawdown", "Nifty fall from its 1-year high", 0.0, -0.30, "How far below the 1-year high Nifty trades."),
    ("Return20", "Nifty 20-day return", 0.05, -0.15, "Speed of the fall: a fast fall creates more panic than a slow one."),
    ("Breadth", "Share of watchlist stocks below their 200-day average", 0.20, 0.90,
     "How widespread the selling is across your 205 stocks."),
    ("Rupee", "USD/INR 20-day change", -0.01, 0.04, "Rupee weakness: foreign money leaving India."),
]
PANIC_ROW0 = 6  # first scale row on the Panic_Meter sheet
SCORE_COLS = [f"{k}_Score" for k, *_ in PANIC_COMPONENTS]


def panic_scale():
    return ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B", mid_type="num", mid_value=50,
                          mid_color="FFEB84", end_type="num", end_value=100, end_color="F8696B")


def build_market_daily(wb, md):
    ws = wb.create_sheet("Market_Daily")
    cols = list(md.columns) + SCORE_COLS + ["Panic"]
    MDC.update({name: L(i + 1) for i, name in enumerate(cols)})
    header(ws, 1, cols, height=45)
    for r, row in enumerate(md.itertuples(index=False), start=2):
        for c, v in enumerate(row, start=1):
            if c == 1:
                v = pd.Timestamp(v).to_pydatetime()
            ws.cell(r, c, nz(v))
        for k, (comp, *_rest) in enumerate(PANIC_COMPONENTS):
            x, row_s = f"{MDC[comp]}{r}", PANIC_ROW0 + k
            calm, fear = f"Panic_Meter!$B${row_s}", f"Panic_Meter!$C${row_s}"
            ws[f"{MDC[SCORE_COLS[k]]}{r}"] = f'=IF({x}="","",MAX(0,MIN(100,({x}-{calm})/({fear}-{calm})*100)))'
        sc = f"{MDC[SCORE_COLS[0]]}{r}:{MDC[SCORE_COLS[-1]]}{r}"
        ws[f"{MDC['Panic']}{r}"] = f'=IF(COUNT({sc})<3,"",AVERAGE({sc}))'
    last = len(md) + 1
    for c, name in enumerate(cols, start=1):
        fmt = ("dd-mmm-yyyy" if name == "Date" else "#,##0.00" if name in ("Nifty", "Brent")
               else "0.00" if name == "VIX" else "0" if name in SCORE_COLS + ["Panic"] else PCT)
        for cell in ws[f"{L(c)}2:{L(c)}{last}"]:
            cell[0].number_format = fmt
    ws.freeze_panes = "B2"
    for i in range(1, len(cols) + 1):
        ws.column_dimensions[L(i)].width = 11
    ws.conditional_formatting.add(f"{MDC['Panic']}2:{MDC['Panic']}{last}", panic_scale())
    return dict(MDC), last


def build_crude(wb, mdc, last_md):
    ws = wb.create_sheet("Crude_Ranges")
    title(ws, "Crude Oil Ranges — at what level of crude move does the market react? (formulas over Market_Daily)",
          "Change the yellow bucket limits to test your own ranges. 'Same 5 days' = during the crude move; "
          "'next' = what followed. Industry columns are equal-weight averages of your watchlist stocks.")
    md = lambda k: f"Market_Daily!${mdc[k]}$2:${mdc[k]}${last_md}"  # noqa: E731
    groups = [k[:-6] for k in mdc if k.endswith("_Same5") and k != "Nifty_Same5"]

    # ---- table A: 5-day Brent change
    ws["A4"] = "A. Brent change over 5 trading days"
    ws["A4"].font = F_BOLD
    heads = ["From", "To", "Days", "Nifty same 5 days", "Nifty next 5 days", "Nifty next 20 days",
             "% Nifty higher after 20 days", "Avg India VIX"] + [f"{g} same 5 days" for g in groups]
    header(ws, 5, heads, height=45)
    buckets = [(-1, -0.15), (-0.15, -0.10), (-0.10, -0.05), (-0.05, -0.02), (-0.02, 0.02),
               (0.02, 0.05), (0.05, 0.10), (0.10, 0.15), (0.15, 5)]
    for i, (lo, hi) in enumerate(buckets):
        r = 6 + i
        for c, v in ((1, lo), (2, hi)):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.number_format, cell.border = F_INPUT, FILL_INPUT, PCT, BOX
        crit = f'{md("Brent_5D")},">="&$A{r},{md("Brent_5D")},"<"&$B{r}'
        ws.cell(r, 3, f"=COUNTIFS({crit})")
        for c, key in ((4, "Nifty_Same5"), (5, "Nifty_Next5"), (6, "Nifty_Next20"), (8, "VIX")):
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({md(key)},{crit}),"")')
        ws.cell(r, 7, f'=IFERROR(COUNTIFS({crit},{md("Nifty_Next20")},">0")/C{r},"")')
        for j, g in enumerate(groups):
            ws.cell(r, 9 + j, f'=IFERROR(AVERAGEIFS({md(g + "_Same5")},{crit}),"")')
    endA = 5 + len(buckets)
    lastc = L(8 + len(groups))
    style_range(ws, f"C6:{lastc}{endA}", PCT)
    style_range(ws, f"C6:C{endA}", NUM0)
    style_range(ws, f"H6:H{endA}", "0.0")
    ws.conditional_formatting.add(f"D6:F{endA}", scale3())
    ws.conditional_formatting.add(f"I6:{lastc}{endA}", scale3())

    # ---- table B: Brent price level
    r0 = endA + 3
    ws.cell(r0 - 1, 1, "B. Brent price level ($ per barrel)").font = F_BOLD
    heads = ["From $", "To $", "Days", "Nifty next 20 days", "% Nifty higher after 20 days", "Avg India VIX"] + \
            [f"{g} next 5 days" for g in groups]
    header(ws, r0, heads, height=45)
    levels = [(0, 40), (40, 60), (60, 80), (80, 100), (100, 120), (120, 200)]
    for i, (lo, hi) in enumerate(levels):
        r = r0 + 1 + i
        for c, v in ((1, lo), (2, hi)):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.number_format, cell.border = F_INPUT, FILL_INPUT, "0", BOX
        crit = f'{md("Brent")},">="&$A{r},{md("Brent")},"<"&$B{r}'
        ws.cell(r, 3, f"=COUNTIFS({crit})")
        ws.cell(r, 4, f'=IFERROR(AVERAGEIFS({md("Nifty_Next20")},{crit}),"")')
        ws.cell(r, 5, f'=IFERROR(COUNTIFS({crit},{md("Nifty_Next20")},">0")/C{r},"")')
        ws.cell(r, 6, f'=IFERROR(AVERAGEIFS({md("VIX")},{crit}),"")')
        for j, g in enumerate(groups):
            ws.cell(r, 7 + j, f'=IFERROR(AVERAGEIFS({md(g + "_Next5")},{crit}),"")')
    endB = r0 + len(levels)
    lastb = L(6 + len(groups))
    style_range(ws, f"C{r0 + 1}:{lastb}{endB}", PCT)
    style_range(ws, f"C{r0 + 1}:C{endB}", NUM0)
    style_range(ws, f"F{r0 + 1}:F{endB}", "0.0")
    ws.conditional_formatting.add(f"D{r0 + 1}:D{endB}", scale3())
    ws.conditional_formatting.add(f"G{r0 + 1}:{lastb}{endB}", scale3())
    ws.cell(endB + 2, 1, "Reading tip: compare the extreme rows with the middle (-2% to +2%) row. The range where "
                         "the numbers clearly change is the level where the market starts reacting to crude.").font = F_NOTE
    widths(ws, {"A": 9, "B": 9})
    for i in range(3, 9 + len(groups)):
        ws.column_dimensions[L(i)].width = 12


def build_fii(wb, mdc, last_md):
    ws = wb.create_sheet("FII_DII")
    title(ws, "FII / DII flows vs Nifty — your daily tracker data (formulas)",
          "Copied from your FII_DII_Daily sheet (Jun-Oct 2026). Paste more history in the yellow columns A-E "
          f"(up to {FII_ROWS} rows); Nifty columns fill in automatically.")
    header(ws, 4, ["Date", "FII net (Rs cr)", "DII net (Rs cr)", "PCR Nifty", "India VIX", "Nifty close",
                   "Nifty move same day", "Nifty move next day", "FII cumulative", "DII cumulative",
                   "Who absorbed?"], height=45)
    rows = []
    if FII_SOURCE.exists():
        fii = pd.read_csv(FII_SOURCE, parse_dates=["Date"])
        rows = [(d.to_pydatetime(), a, b, c, e) for d, a, b, c, e in fii.itertuples(index=False)]
    md = lambda k: f"Market_Daily!${mdc[k]}$2:${mdc[k]}${last_md}"  # noqa: E731
    end = 4 + FII_ROWS
    for i in range(FII_ROWS):
        r = 5 + i
        if i < len(rows):
            d, fii, dii, pcr, vix = (nz(v) for v in rows[i])
            for c, v in ((1, d), (2, fii), (3, dii), (4, pcr), (5, vix)):
                ws.cell(r, c, v)
        for c in range(1, 6):
            ws.cell(r, c).fill, ws.cell(r, c).font = FILL_INPUT, F_INPUT
        ws.cell(r, 6, f'=IF(A{r}="","",IFERROR(INDEX({md("Nifty")},MATCH(A{r},{md("Date")},0)),""))')
        ws.cell(r, 7, f'=IF(A{r}="","",IFERROR(INDEX({md("Nifty_Ret1")},MATCH(A{r},{md("Date")},0)),""))')
        ws.cell(r, 8, f'=IF(A{r}="","",IFERROR(INDEX({md("Nifty_Next1")},MATCH(A{r},{md("Date")},0)),""))')
        ws.cell(r, 9, f'=IF(B{r}="","",SUM($B$5:B{r}))')
        ws.cell(r, 10, f'=IF(C{r}="","",SUM($C$5:C{r}))')
        ws.cell(r, 11, f'=IF(OR(B{r}="",C{r}=""),"",IF(AND(B{r}<0,C{r}>0),"DII bought FII selling",'
                       f'IF(AND(B{r}>0,C{r}<0),"FII bought DII selling",IF(B{r}>0,"Both buying","Both selling"))))')
    for col, fmt in (("A", "dd-mmm-yyyy"), ("B", CR), ("C", CR), ("D", "0.00"), ("E", "0.00"), ("F", "#,##0.00"),
                     ("G", PCT), ("H", PCT), ("I", CR), ("J", CR)):
        for c in ws[f"{col}5:{col}{end}"]:
            c[0].number_format = fmt
    ws.freeze_panes = "B5"

    # ---- summary block
    rng = lambda c: f"${c}$5:${c}${end}"  # noqa: E731
    th = S["fii"]
    items = [
        ("Days with data", f'=COUNT({rng("B")})', "0"),
        ("Total FII net (Rs cr)", f'=SUM({rng("B")})', CR),
        ("Total DII net (Rs cr)", f'=SUM({rng("C")})', CR),
        ("Correlation: FII flow vs Nifty same day", f'=IFERROR(CORREL({rng("B")},{rng("G")}),"")', "0.00"),
        ("Correlation: FII flow vs Nifty NEXT day", f'=IFERROR(CORREL({rng("B")},{rng("H")}),"")', "0.00"),
        ("Correlation: DII flow vs Nifty same day", f'=IFERROR(CORREL({rng("C")},{rng("G")}),"")', "0.00"),
        ("Nifty avg move on BIG FII BUY days", f'=IFERROR(AVERAGEIFS({rng("G")},{rng("B")},">="&{th}),"")', PCT),
        ("Nifty avg move on BIG FII SELL days", f'=IFERROR(AVERAGEIFS({rng("G")},{rng("B")},"<="&-{th}),"")', PCT),
        ("Nifty avg NEXT-day move after BIG FII SELL", f'=IFERROR(AVERAGEIFS({rng("H")},{rng("B")},"<="&-{th}),"")', PCT),
        ("Nifty avg move on BIG DII BUY days", f'=IFERROR(AVERAGEIFS({rng("G")},{rng("C")},">="&{th}),"")', PCT),
        ("Days DII absorbed FII selling", f'=COUNTIF({rng("K")},"DII bought FII selling")', "0"),
        ("Days FII absorbed DII selling", f'=COUNTIF({rng("K")},"FII bought DII selling")', "0"),
        ("Nifty avg move when DII absorbed FII selling",
         f'=IFERROR(AVERAGEIFS({rng("G")},{rng("K")},"DII bought FII selling"),"")', PCT),
    ]
    ws["M4"] = "Summary"
    ws["M4"].font, ws["M4"].fill = F_HEAD, FILL_HEAD
    ws["N4"].fill = FILL_HEAD
    for i, (label, f, fmt) in enumerate(items):
        r = 5 + i
        ws.cell(r, 13, label).font = F_BASE
        c = ws.cell(r, 14, f)
        c.font, c.number_format, c.border = F_BOLD, fmt, BOX
    ws.cell(5 + len(items) + 1, 13, "Long-term FII/DII history (2007 onwards) is needed to link flows to each event. "
                                    "Download it from NSE/NSDL and paste here.").font = F_NOTE
    widths(ws, {"A": 12, "B": 12, "C": 12, "D": 9, "E": 9, "F": 11, "G": 11, "H": 11, "I": 12, "J": 12,
                "K": 22, "L": 3, "M": 44, "N": 12})


# ---------------------------------------------------------------- target adjuster
def build_adjuster(wb, last_se, end_sc, cat_list, example):
    ws = wb.create_sheet("Target_Adjuster")
    title(ws, "Target Adjuster — adjust your mathematical targets for the emotion of an upcoming event",
          "Enter the yellow cells. Your software's targets stay the base; this sheet adds the historical event effect.")
    stk, cat = se_rng("A", last_se), se_rng("D", last_se)
    sc = lambda c: f"Stock_Scorecard!${c}$5:${c}${end_sc}"  # noqa: E731
    cs = lambda c: f"Category_Summary!${c}$5:${c}${cat_list.split('$')[-1]}"  # noqa: E731
    inputs = [("Stock (NSE code)", example["stock"], None), ("Upcoming event type", example["category"], None),
              ("Current market price (CMP)", example["cmp"], RS),
              ("Base short-term target (from your model)", example["st"], RS),
              ("Base long-term target (from your model)", example["lt"], RS)]
    header(ws, 3, ["INPUTS", "Value"])
    for i, (label, v, fmt) in enumerate(inputs):
        r = 4 + i
        ws.cell(r, 1, label).font = F_BASE
        c = ws.cell(r, 2, v)
        c.font, c.fill, c.border = F_INPUT, FILL_INPUT, BOX
        if fmt:
            c.number_format = fmt
    ws["C6"] = "Example values (last close and model-style targets): replace with your own."
    ws["C6"].font = F_NOTE
    dv1 = DataValidation(type="list", formula1=f"={sc('A')}", allow_blank=False)
    dv2 = DataValidation(type="list", formula1=f"={cat_list}", allow_blank=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add("B4")
    dv2.add("B5")

    crit = f"{stk},$B$4,{cat},$B$5"
    rows = [
        ("HISTORY — this stock in this event type", None, None),
        ("Times measured", f"=COUNTIFS({crit})", "0"),
        ("Avg event-day move", f'=IFERROR(AVERAGEIFS({se_rng("H", last_se)},{crit}),"")', PCT),
        ("Avg total move (20 days)", f'=IFERROR(AVERAGEIFS({se_rng("K", last_se)},{crit}),"")', PCT),
        ("Avg worst fall (20 days)", f'=IFERROR(AVERAGEIFS({se_rng("N", last_se)},{crit}),"")', PCT),
        ("Avg days to recover", f'=IFERROR(AVERAGEIFS({se_rng("P", last_se)},{crit}),"")', "0"),
        ("R:R when bought on event day", f'=IFERROR(AVERAGEIFS({se_rng("R", last_se)},{crit})/AVERAGEIFS({se_rng("S", last_se)},{crit}),"")', "0.00"),
        ("MARKET — Nifty in this event type", None, None),
        ("Nifty avg total move (20 days)", f'=IFERROR(INDEX({cs("F")},MATCH($B$5,{cs("A")},0)),"")', PCT),
        ("Nifty avg worst fall (20 days)", f'=IFERROR(INDEX({cs("I")},MATCH($B$5,{cs("A")},0)),"")', PCT),
        ("Direction points", f'=IFERROR(INDEX({cs("W")},MATCH($B$5,{cs("A")},0)),"")', "0"),
        ("Risk points", f'=IFERROR(INDEX({cs("X")},MATCH($B$5,{cs("A")},0)),"")', "0"),
        ("Best entry (history)", f'=IFERROR(INDEX({cs("V")},MATCH($B$5,{cs("A")},0)),"")', None),
        ("Stock Event Beta", f'=IFERROR(INDEX({sc("O")},MATCH($B$4,{sc("A")},0)),"")', "0.00"),
        ("RESULT", None, None),
        ("Method used", f'=IF(B10>={S["minev"]},"Stock\'s own history","Event Beta x Nifty (too few stock events)")', None),
        ("Expected event move", f'=IFERROR(IF(B10>={S["minev"]},B12,B22*B17),0)', PCT),
        ("Expected worst dip", f'=IFERROR(IF(B10>={S["minev"]},B13,B22*B18),0)', PCT),
        ("Adjusted short-term target", f'=B7*(1+B25*{S["stw"]})', RS),
        ("Adjusted long-term target", f'=B8*(1+B25*{S["ltw"]})', RS),
        ("Suggested stop-loss", f'=B6*(1+MIN(B26,0)*{S["slm"]})', RS),
        ("R:R from CMP (target gain / stop-loss risk)", '=IFERROR((B27-B6)/(B6-B29),"")', "0.00"),
        ("Short-term target change vs base", '=IFERROR(B27/B7-1,"")', PCT),
    ]
    for i, (label, f, fmt) in enumerate(rows):
        r = 9 + i
        if f is None:
            header(ws, r, [label, ""], fill=FILL_SUB, font=F_BOLD, height=18)
            continue
        ws.cell(r, 1, label).font = F_BASE
        c = ws.cell(r, 2, f)
        c.font, c.border = F_BOLD, BOX
        if fmt:
            c.number_format = fmt
        if r >= 24:
            c.fill = FILL_OUT
    ws["C25"] = "Positive = event historically lifted this stock; negative = it pulled it down."
    ws["C29"] = "Based on the typical worst dip after this event type (Settings: stop-loss buffer)."
    for a in ("C25", "C29"):
        ws[a].font = F_NOTE
    widths(ws, {"A": 46, "B": 22, "C": 80})



# ---------------------------------------------------------------- panic meter
PANIC_ZONES = [(0, 20, "Complacent / greedy"), (20, 40, "Calm"), (40, 60, "Worried"),
               (60, 80, "Fear"), (80, 101, "Panic")]


def build_panic(wb, mdc, last_md, last_ni):
    ws = wb.create_sheet("Panic_Meter")
    title(ws, "Panic Meter — how frightened is the market (0 = calm, 100 = extreme panic), and what followed?",
          "Each component is scored 0-100 between its calm and fear values (yellow, editable); the Panic score is "
          "their average. Fixed scales are used, so no hindsight goes into the score.")
    header(ws, 5, ["Component", "Calm value (score 0)", "Fear value (score 100)", "Meaning"])
    for k, (comp, label, calm, fear, meaning) in enumerate(PANIC_COMPONENTS):
        r = PANIC_ROW0 + k
        ws.cell(r, 1, label).font = F_BASE
        fmt = "0" if comp == "VIX" else PCT
        for c, v in ((2, calm), (3, fear)):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.number_format, cell.border = F_INPUT, FILL_INPUT, fmt, BOX
        ws.cell(r, 4, meaning).font = F_NOTE

    md = lambda k: f"Market_Daily!${mdc[k]}$2:${mdc[k]}${last_md}"  # noqa: E731
    # ---- latest reading
    r0 = PANIC_ROW0 + len(PANIC_COMPONENTS) + 1
    header(ws, r0, ["LATEST READING", "Value"], fill=FILL_SUB, font=F_BOLD, height=18)
    ws.cell(r0 + 1, 1, "Latest date in data")
    ws.cell(r0 + 1, 2, f"=MAX({md('Date')})").number_format = "dd-mmm-yyyy"
    ws.cell(r0 + 2, 1, "Panic score")
    ws.cell(r0 + 2, 2, f"=INDEX({md('Panic')},MATCH(B{r0 + 1},{md('Date')},0))").number_format = "0"
    ws.cell(r0 + 3, 1, "Zone")
    z0 = r0 + 7
    zr = f"{z0 + 1}:$C${z0 + len(PANIC_ZONES)}"
    ws.cell(r0 + 3, 2, f'=IFERROR(INDEX($C${zr},MATCH(B{r0 + 2},$A${z0 + 1}:$A${z0 + len(PANIC_ZONES)},1)),"")')
    for rr in range(r0 + 1, r0 + 4):
        ws.cell(rr, 1).font, ws.cell(rr, 2).font, ws.cell(rr, 2).fill = F_BASE, F_BOLD, FILL_OUT
    ws.conditional_formatting.add(f"B{r0 + 2}", panic_scale())

    # ---- zones: what followed
    ws.cell(z0 - 1, 1, "A. Daily history since 2007: what Nifty did next at each panic level").font = F_BOLD
    header(ws, z0, ["From", "To", "Zone", "Days", "% of all days", "Nifty next 20 days", "Nifty next 60 days",
                    "Nifty next 1 year", "% higher after 1 year", "Worst next 60 days"], height=45)
    for i, (lo, hi, name) in enumerate(PANIC_ZONES):
        r = z0 + 1 + i
        for c, v in ((1, lo), (2, hi)):
            cell = ws.cell(r, c, v)
            cell.font, cell.fill, cell.border = F_INPUT, FILL_INPUT, BOX
        ws.cell(r, 3, name)
        crit = f'{md("Panic")},">="&$A{r},{md("Panic")},"<"&$B{r}'
        ws.cell(r, 4, f"=COUNTIFS({crit})")
        ws.cell(r, 5, f'=IFERROR(D{r}/COUNT({md("Panic")}),"")')
        for c, key in ((6, "Nifty_Next20"), (7, "Nifty_Next60"), (8, "Nifty_Next250")):
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({md(key)},{crit}),"")')
        ws.cell(r, 9, f'=IFERROR(COUNTIFS({crit},{md("Nifty_Next250")},">0")/COUNTIFS({crit},{md("Nifty_Next250")},"<>"),"")')
        ws.cell(r, 10, f'=IFERROR(_xlfn.MINIFS({md("Nifty_Next60")},{crit}),"")')
    zend = z0 + len(PANIC_ZONES)
    style_range(ws, f"C{z0 + 1}:J{zend}")
    for col in "EFGHIJ":
        style_range(ws, f"{col}{z0 + 1}:{col}{zend}", PCT)
    style_range(ws, f"D{z0 + 1}:D{zend}", NUM0)
    ws.conditional_formatting.add(f"F{z0 + 1}:H{zend}", scale3())

    # ---- events: panic vs outcome
    e0 = zend + 4
    ws.cell(e0 - 1, 1, "B. Events: did the panic level match what the market did next?").font = F_BOLD
    header(ws, e0, ["From", "To", "Zone (panic peak within 20 days of event)", "Events", "Avg event-day move",
                    "Avg worst fall (20 days)", "Avg move next 60 days", "% higher after 60 days",
                    "Avg days to recover"], height=45)
    pk = ni_rng("PanicPeak", last_ni)
    for i, (lo, hi, name) in enumerate(PANIC_ZONES):
        r = e0 + 1 + i
        ws.cell(r, 1, f"=A{z0 + 1 + i}")
        ws.cell(r, 2, f"=B{z0 + 1 + i}")
        ws.cell(r, 3, name)
        crit = f'{pk},">="&$A{r},{pk},"<"&$B{r}'
        ws.cell(r, 4, f"=COUNTIFS({crit})")
        for c, key in ((5, "Day0"), (6, "MaxFall20"), (7, "Post60"), (9, "Recovery_Days")):
            ws.cell(r, c, f'=IFERROR(AVERAGEIFS({ni_rng(key, last_ni)},{crit}),"")')
        p60 = ni_rng("Post60", last_ni)
        ws.cell(r, 8, f'=IFERROR(COUNTIFS({crit},{p60},">0")/COUNTIFS({crit},{p60},"<>"),"")')
    eend = e0 + len(PANIC_ZONES)
    style_range(ws, f"A{e0 + 1}:I{eend}")
    for col in "EFGH":
        style_range(ws, f"{col}{e0 + 1}:{col}{eend}", PCT)
    style_range(ws, f"I{e0 + 1}:I{eend}", NUM0)
    ws.conditional_formatting.add(f"G{e0 + 1}:G{eend}", scale3())
    ws.cell(eend + 2, 1, "How to read: if the 'Panic' and 'Fear' rows show the best returns afterwards, extreme fear has "
                         "historically been a buying opportunity: the crowd (and the forecasters) were too gloomy.").font = F_NOTE
    widths(ws, {"A": 46, "B": 14, "C": 34, "D": 90})
    for col in "EFGHIJ":
        ws.column_dimensions[col].width = 13


# ---------------------------------------------------------------- crisis periods
CRISIS_SECTOR_ROWS = {}


def build_crisis(wb, cp, cs, watch, mdc, last_md):
    ws = wb.create_sheet("Crisis_Periods")
    title(ws, "Crisis Periods — recessions and bear markets measured from peak to bottom to full recovery",
          "Peak = highest close before the bottom; Recovery = first close back at the peak. Days are trading days "
          "(about 250 a year). Sensex is used before Sep-2007.")
    heads = ["Crisis", "Index", "Peak date", "Peak level", "Bottom date", "Bottom level", "Fall peak to bottom",
             "Trading days falling", "Recovery date", "Trading days bottom to recovery", "Total months peak to recovery",
             "India VIX peak", "Panic Meter peak", "Return 1 year after bottom",
             "PHASE 1: first 10% fall — date", "PHASE 1: further fall after buying", "PHASE 1: 1-year return",
             "PHASE 2: bottom — 1-year return", "PHASE 3: 20% off bottom — date", "PHASE 3: 1-year return",
             "Practical choice: early vs wait (bottom cannot be timed)", "What caused it", "Status"]
    header(ws, 4, heads, height=60)
    pc, dc = f"Market_Daily!${mdc['Panic']}$2:${mdc['Panic']}${last_md}", f"Market_Daily!$A$2:$A${last_md}"
    dt = lambda v: pd.Timestamp(v).to_pydatetime() if isinstance(v, str) else None  # noqa: E731
    for r, x in enumerate(cp.itertuples(index=False), start=5):
        vals = {1: x.Crisis, 2: x.Index, 3: dt(x.Peak_Date), 4: x.Peak, 5: dt(x.Trough_Date), 6: x.Trough,
                8: x.Days_Down, 9: dt(x.Recovery_Date) or "Not yet", 10: nz(x.Days_To_Recover), 12: nz(x.VIX_Peak),
                14: nz(x.Ret1Y_After_Trough), 15: dt(x.Buy10_Date), 16: nz(x.Buy10_FurtherFall),
                17: nz(x.Buy10_Ret1Y), 18: nz(x.BuyTrough_Ret1Y), 19: dt(x.Buy20Up_Date),
                20: nz(x.Buy20Up_Ret1Y), 22: x.Cause, 23: x.Status}
        for c, v in vals.items():
            ws.cell(r, c, v)
        ws.cell(r, 7, f"=F{r}/D{r}-1")
        ws.cell(r, 11, f'=IF(ISNUMBER(I{r}),(I{r}-C{r})/30.4,"")')
        ws.cell(r, 13, f'=IF(C{r}<MIN({dc}),"",IFERROR(_xlfn.MAXIFS({pc},{dc},">="&C{r},{dc},"<="&E{r}+30),""))')
        ws.cell(r, 21, f'=IF(OR(Q{r}="",T{r}=""),"",IF(Q{r}>=T{r},"Buying early (first 10% fall) paid more",'
                       f'"Waiting for the 20% bounce paid more"))')
    end = 4 + len(cp)
    avg_r = end + 1
    ws.cell(avg_r, 1, "AVERAGE of recovered crises").font = F_BOLD
    for col in "GHJKMNPQRT":
        ws[f"{col}{avg_r}"] = f'=IFERROR(AVERAGEIF($W$5:$W${end},"Recovered",{col}5:{col}{end}),"")'
    style_range(ws, f"A5:W{avg_r}")
    for col in "CEIOS":
        style_range(ws, f"{col}5:{col}{end}", "dd-mmm-yyyy")
    for col in "DF":
        style_range(ws, f"{col}5:{col}{end}", "#,##0")
    for col in "GNPQRT":
        style_range(ws, f"{col}5:{col}{avg_r}", PCT)
    for col in "HJ":
        style_range(ws, f"{col}5:{col}{avg_r}", NUM0)
    style_range(ws, f"K5:K{avg_r}", NUM1)
    style_range(ws, f"L5:M{end}", "0")
    for c in ws[f"A{avg_r}:W{avg_r}"][0]:
        c.font, c.fill = F_BOLD, FILL_SUB
    for r in range(5, end + 1):
        ws.cell(r, 1).font = F_BOLD
        ws.cell(r, 22).alignment = Alignment(wrap_text=True, vertical="top")
    ws.conditional_formatting.add(f"M5:M{end}", panic_scale())
    ws.conditional_formatting.add(f"N5:N{end}", scale3())
    ws.conditional_formatting.add(f"Q5:T{end}", scale3())
    ws.cell(avg_r + 2, 1, "Lesson: buying at the exact bottom is impossible to time. Compare Phase 1 (buy early) with "
                          "Phase 3 (wait for a 20% bounce): Phase 3 gives up some gain but avoids the further fall "
                          "shown in column P.").font = F_NOTE
    ws.conditional_formatting.add(f"W5:W{end}", CellIsRule(operator="equal", formula=['"Recovered"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"W5:W{end}", CellIsRule(operator="notEqual", formula=['"Recovered"'], fill=RED_FILL))

    # ---- where are we now (last crisis row = the current one)
    n_ = f"Market_Daily!${mdc['Nifty']}$2:${mdc['Nifty']}${last_md}"
    d_ = f"Market_Daily!$A$2:$A${last_md}"
    b_ = f"Market_Daily!${mdc['Breadth']}$2:${mdc['Breadth']}${last_md}"
    k = end
    w0 = avg_r + 5
    ws.cell(w0 - 1, 1, "WHERE ARE WE NOW? The current crisis compared with history (formulas)").font = Font(
        name=FONT, size=12, bold=True, color="C00000")
    header(ws, w0, ["Measure", "Value", "Compared with past crises"], height=20)
    now_rows = [
        ("Latest date", f"=MAX({d_})", "dd-mmm-yyyy", ""),
        ("Latest Nifty", f"=INDEX({n_},COUNT({n_}))", "#,##0", ""),
        ("Peak of this crisis", f"=D{k}", "#,##0", ""),
        ("Lowest close so far", f"=F{k}", "#,##0", ""),
        ("Fall at the low so far", f"=G{k}", PCT, f'=IFERROR("Average past fall: "&TEXT(G{avg_r},"0%"),"")'),
        ("Nifty now vs peak", f"=B{w0 + 2}/B{w0 + 3}-1", PCT, ""),
        ("Nifty now vs the low", f"=B{w0 + 2}/B{w0 + 4}-1", PCT, "20% above the low has confirmed recoveries in the past"),
        ("Trading days since the peak", f'=COUNTIFS({d_},">"&C{k})', "0",
         f'=IFERROR("Average days falling: "&TEXT(H{avg_r},"0"),"")'),
        ("Share of past crises that fell deeper", f'=IFERROR(COUNTIFS($G$5:$G${k - 1},"<"&G{k})/COUNT($G$5:$G${k - 1}),"")',
         PCT, ""),
        ("Panic Meter peak in this crisis", f"=M{k}", "0", f'=IFERROR("Average panic peak at past bottoms: "&TEXT(M{avg_r},"0"),"")'),
        ("Panic Meter now", "=Panic_Meter!$B$14", "0", "Fear 60-80, Panic 80+ have marked past bottoms"),
        ("Nifty vs its 200-day average", f"=B{w0 + 2}/AVERAGE(INDEX({n_},COUNT({n_})-199):INDEX({n_},COUNT({n_})))-1",
         PCT, "Below zero = long-term trend still down"),
        ("Watchlist stocks below 200-day average", f"=INDEX({b_},COUNT({n_}))", PCT, ""),
        ("PHASE NOW", f'=IF(B{w0 + 2}>=B{w0 + 3},"Recovered: back at the peak",IF(B{w0 + 2}<=B{w0 + 4}*1.03,'
                      f'"Testing the low (within 3%)",IF(B{w0 + 2}>=B{w0 + 4}*1.2,"Recovery confirmed (20%+ above the low)",'
                      f'"Bounce off the low, recovery not confirmed")))', None, ""),
    ]
    for i, (label, f, fmt, cmp_) in enumerate(now_rows):
        r = w0 + 1 + i
        ws.cell(r, 1, label).font = F_BOLD if label == "PHASE NOW" else F_BASE
        c = ws.cell(r, 2, f)
        c.font, c.border, c.fill = F_BOLD, BOX, FILL_OUT
        if fmt:
            c.number_format = fmt
        if cmp_:
            ws.cell(r, 3, cmp_).font = F_NOTE
    ws.conditional_formatting.add(f"B{w0 + 10}:B{w0 + 11}", panic_scale())
    ws.freeze_panes = "B5"
    widths(ws, {"A": 40, "V": 70, "W": 22})
    for i in range(2, 22):
        ws.column_dimensions[L(i)].width = 12

    # ---- raw stock rows
    meta = watch.set_index("NSE_Code")
    wd = wb.create_sheet("Crisis_Stock_Data")
    cols = ["Stock", "Crisis", "Industry", "Cap", "Fall (index peak to bottom)", "Worst fall", "Vs index",
            "Trading days to recover", "Recovered? (1=yes)", "Return 1 year after bottom"]
    header(wd, 1, cols, height=45)
    for r, x in enumerate(cs.itertuples(index=False), start=2):
        for c, v in enumerate([x.Stock, x.Crisis, meta["Industry"].get(x.Stock), meta["Cap"].get(x.Stock),
                               x.Fall_Peak_To_Trough, x.Max_Fall, x.Vs_Index, nz(x.Days_To_Recover),
                               x.Recovered, nz(x.Ret1Y_After_Trough)], start=1):
            wd.cell(r, c, v)
    last_cs = len(cs) + 1
    for col in "EFGJ":
        for c in wd[f"{col}2:{col}{last_cs}"]:
            c[0].number_format = PCT
    wd.freeze_panes = "C2"
    wd.auto_filter.ref = f"A1:J{last_cs}"
    widths(wd, {"A": 13, "B": 40, "C": 18})

    # ---- stock scorecard
    rng = lambda col: f"Crisis_Stock_Data!${col}$2:${col}${last_cs}"  # noqa: E731
    wc = wb.create_sheet("Crisis_Scorecard")
    title(wc, "Crisis Scorecard — how each watchlist stock behaved in recessions and bear markets (formulas)",
          "Defence = fell less than the index. Recovery leader = best 1-year return after the market bottom. "
          "Ranked only with enough crises measured (Settings).")
    header(wc, 4, ["Stock", "Name", "Industry", "Cap", "Crises measured", "Avg fall (peak to bottom)",
                   "Avg worst fall", "Avg vs index", "% of crises fully recovered", "Avg trading days to recover",
                   "Avg 1-year return after bottom", "Defence score", "Defence rank (1 = best)",
                   "Recovery score", "Recovery rank (1 = best)"], height=60)
    n = len(watch)
    send = 4 + n
    for r, w in enumerate(watch.itertuples(index=False), start=5):
        wc.cell(r, 1, w.NSE_Code)
        wc.cell(r, 2, w.Name)
        wc.cell(r, 3, w.Industry)
        wc.cell(r, 4, w.Cap)
        wc.cell(r, 5, f'=COUNTIF({rng("A")},$A{r})')
        for c, col in ((6, "E"), (7, "F"), (8, "G"), (9, "I"), (10, "H"), (11, "J")):
            wc.cell(r, c, f'=IFERROR(AVERAGEIFS({rng(col)},{rng("A")},$A{r}),"")')
        ok = f'OR(E{r}<{S["mincr"]},H{r}="",K{r}="")'
        wc.cell(r, 12, f'=IF({ok},"",RANK(H{r},$H$5:$H${send},0)+ROW()/100000)')
        wc.cell(r, 13, f'=IF(L{r}="","",RANK(L{r},$L$5:$L${send},1))')
        wc.cell(r, 14, f'=IF({ok},"",RANK(K{r},$K$5:$K${send},0)+ROW()/100000)')
        wc.cell(r, 15, f'=IF(N{r}="","",RANK(N{r},$N$5:$N${send},1))')
    style_range(wc, f"A5:O{send}")
    for r in range(5, send + 1):
        wc.cell(r, 1).font = F_BOLD
    for col in "FGHIK":
        style_range(wc, f"{col}5:{col}{send}", PCT)
    style_range(wc, f"J5:J{send}", NUM0)
    for col in "LN":
        style_range(wc, f"{col}5:{col}{send}", "0.0")
    for col in "HK":
        wc.conditional_formatting.add(f"{col}5:{col}{send}", scale3())
    wc.freeze_panes = "B5"
    wc.auto_filter.ref = f"A4:O{send}"
    widths(wc, {"A": 13, "B": 20, "C": 18, "D": 10})
    for col in "EFGHIJKLMNO":
        wc.column_dimensions[col].width = 12

    # ---- sector x crisis
    wsx = wb.create_sheet("Crisis_Sectors")
    title(wsx, "Crisis Sectors — average fall of watchlist stocks in each industry, by crisis (formulas)",
          "Read across a row to see whether an industry is consistently defensive; the last column shows "
          "the average 1-year bounce after the market bottom.")
    crises = list(cp["Crisis"])
    header(wsx, 4, ["Industry"] + crises + ["Avg fall, all crises", "Avg 1-year return after bottom"], height=75)
    industries = sorted(watch["Industry"].dropna().unique())
    for r, ind in enumerate(industries, start=5):
        wsx.cell(r, 1, ind).font = F_BOLD
        for j, name in enumerate(crises):
            wsx.cell(r, 2 + j, f'=IFERROR(AVERAGEIFS({rng("E")},{rng("C")},$A{r},{rng("B")},{L(2 + j)}$4),"")')
        k = 2 + len(crises)
        wsx.cell(r, k, f'=IFERROR(AVERAGEIFS({rng("E")},{rng("C")},$A{r}),"")')
        wsx.cell(r, k + 1, f'=IFERROR(AVERAGEIFS({rng("J")},{rng("C")},$A{r}),"")')
    iend = 4 + len(industries)
    k = 2 + len(crises)
    style_range(wsx, f"B5:{L(k + 1)}{iend}", PCT)
    wsx.conditional_formatting.add(f"B5:{L(k)}{iend}", scale3())
    wsx.conditional_formatting.add(f"{L(k + 1)}5:{L(k + 1)}{iend}", scale3())
    # stocks per industry (best/worst lists use only industries with at least 3 watchlist stocks)
    cnt_col = L(k + 2)
    header(wsx, 4, ["Watchlist stocks"], col=k + 2, height=75)
    for r, ind in enumerate(industries, start=5):
        wsx.cell(r, k + 2, f'=COUNTIF(Crisis_Scorecard!$C$5:$C${send},$A{r})').border = BOX
    cnt = f"${cnt_col}$5:${cnt_col}${iend}"
    # best / worst industry per crisis (fall)
    ind_rng = f"$A$5:$A${iend}"
    br = iend + 1
    for off, (lab, fn) in enumerate((("Held up best in the fall", "MAX"), ("Hit hardest in the fall", "MIN"))):
        r = br + off
        wsx.cell(r, 1, lab).font = F_BOLD
        for j in range(len(crises)):
            col = L(2 + j)
            c = wsx.cell(r, 2 + j, f'=IFERROR(INDEX({ind_rng},MATCH(_xlfn.{fn}IFS({col}5:{col}{iend},{cnt},">=3"),{col}5:{col}{iend},0)),"")')
            c.font, c.fill, c.alignment = F_BOLD, FILL_SUB, Alignment(wrap_text=True)
    # 1-year after bottom matrix
    y0 = br + 4
    wsx.cell(y0 - 1, 1, "Average 1-year return after the market bottom, by industry and crisis (where the money went in the recovery)").font = F_BOLD
    header(wsx, y0, ["Industry"] + crises, height=75)
    for i, ind in enumerate(industries):
        r = y0 + 1 + i
        wsx.cell(r, 1, ind).font = F_BOLD
        for j in range(len(crises)):
            wsx.cell(r, 2 + j, f'=IFERROR(AVERAGEIFS({rng("J")},{rng("C")},$A{r},{rng("B")},{L(2 + j)}${y0}),"")')
    yend = y0 + len(industries)
    style_range(wsx, f"B{y0 + 1}:{L(1 + len(crises))}{yend}", PCT)
    wsx.conditional_formatting.add(f"B{y0 + 1}:{L(1 + len(crises))}{yend}", scale3())
    ind_rng2 = f"$A${y0 + 1}:$A${yend}"
    for off, (lab, fn) in enumerate((("Recovery leader", "MAX"), ("Recovery laggard", "MIN"))):
        r = yend + 1 + off
        wsx.cell(r, 1, lab).font = F_BOLD
        for j in range(len(crises)):
            col = L(2 + j)
            c = wsx.cell(r, 2 + j, f'=IFERROR(INDEX({ind_rng2},MATCH(_xlfn.{fn}IFS({col}{y0 + 1}:{col}{yend},{cnt},">=3"),{col}{y0 + 1}:{col}{yend},0)),"")')
            c.font, c.fill, c.alignment = F_BOLD, FILL_SUB, Alignment(wrap_text=True)
    wsx.freeze_panes = "B5"
    widths(wsx, {"A": 22})
    for i in range(2, k + 2):
        wsx.column_dimensions[L(i)].width = 12
    CRISIS_SECTOR_ROWS.update(best=br, worst=br + 1, lead=yend + 1, lag=yend + 2, n=len(crises))
    return send


def add_crisis_top_lists(wb, send, top=20):
    ws = wb["Top_Lists"]
    sc = lambda c: f"Crisis_Scorecard!${c}$5:${c}${send}"  # noqa: E731
    base = ws.max_row + 3
    blocks = [("CRISIS DEFENDERS — fell least versus the index in recessions and bear markets", "L", base),
              ("RECOVERY LEADERS — best 1-year return after the market bottom", "N", base + top + 4)]
    for label, score_col, r0 in blocks:
        ws.cell(r0 - 1, 1, label).font = F_BOLD
        header(ws, r0, ["Rank", "Stock", "Name", "Industry", "Cap", "Crises measured", "Avg fall",
                        "Avg vs index", "% recovered", "Avg days to recover", "Avg 1-year after bottom"], height=45)
        for k in range(1, top + 1):
            r = r0 + k
            ws.cell(r, 1, k)
            match = f"MATCH(SMALL({sc(score_col)},$A{r}),{sc(score_col)},0)"
            for c, col in zip(range(2, 12), "ABCDEFHIJK"):
                ws.cell(r, c, f'=IFERROR(INDEX({sc(col)},{match}),"")')
        style_range(ws, f"A{r0 + 1}:K{r0 + top}")
        for col in "GHIK":
            style_range(ws, f"{col}{r0 + 1}:{col}{r0 + top}", PCT)
        style_range(ws, f"J{r0 + 1}:J{r0 + top}", NUM0)


# ---------------------------------------------------------------- global recovery
def build_global(wb, gc, crises):
    wd = wb.create_sheet("Global_Crisis_Data")
    cols = ["Country", "Index", "Crisis", "Peak date", "Bottom date", "Fall", "Trading days falling",
            "Recovery date", "Trading days to recover", "Recovered? (1=yes)", "1-year return after bottom",
            "Today vs that crisis peak"]
    header(wd, 1, cols, height=45)
    dt = lambda v: pd.Timestamp(v).to_pydatetime() if isinstance(v, str) else "Not yet"  # noqa: E731
    for r, x in enumerate(gc.itertuples(index=False), start=2):
        for c, v in enumerate([x.Country, x.Index, x.Crisis, dt(x.Peak_Date), dt(x.Trough_Date), x.Fall, x.Days_Down,
                               dt(x.Recovery_Date), nz(x.Days_To_Recover), x.Recovered, nz(x.Ret1Y_After_Trough),
                               x.Now_Vs_Peak], start=1):
            wd.cell(r, c, v)
    last = len(gc) + 1
    for col, fmt in (("D", "dd-mmm-yyyy"), ("E", "dd-mmm-yyyy"), ("H", "dd-mmm-yyyy"), ("F", PCT), ("K", PCT), ("L", PCT)):
        for c in wd[f"{col}2:{col}{last}"]:
            c[0].number_format = fmt
    wd.freeze_panes = "B2"
    wd.auto_filter.ref = f"A1:L{last}"
    widths(wd, {"A": 22, "B": 24, "C": 40})

    rng = lambda c: f"Global_Crisis_Data!${c}$2:${c}${last}"  # noqa: E731
    ws = wb.create_sheet("Global_Recovery")
    title(ws, "Global Recovery — how each country's stock market fell and recovered in every crisis (formulas)",
          "Local-currency price indices (no dividends). Peak and bottom are found inside each crisis window. "
          "'Not yet' = the index has never regained that crisis peak.")
    countries = list(dict.fromkeys(gc["Country"]))
    nc = len(crises)

    def matrix(r0, label, col, fmt, not_yet=False):
        ws.cell(r0 - 1, 1, label).font = F_BOLD
        header(ws, r0, ["Country"] + crises, height=75)
        for i, ctry in enumerate(countries):
            r = r0 + 1 + i
            ws.cell(r, 1, ctry).font = F_BOLD
            for j in range(nc):
                h = f"{L(2 + j)}${r0}"
                f = f'AVERAGEIFS({rng(col)},{rng("A")},$A{r},{rng("C")},{h})'
                if not_yet:
                    f = f'IF(COUNTIFS({rng("A")},$A{r},{rng("C")},{h},{rng("J")},0)>0,"Not yet",{f})'
                ws.cell(r, 2 + j, f'=IFERROR({f},"")')
        end = r0 + len(countries)
        style_range(ws, f"B{r0 + 1}:{L(1 + nc)}{end}", fmt)
        return end

    a0 = 5
    aend = matrix(a0, "A. Fall from peak to bottom", "F", PCT)
    ws.conditional_formatting.add(f"B{a0 + 1}:{L(1 + nc)}{aend}", scale3())
    # summary to the right of A
    sc = 3 + nc
    header(ws, a0, ["Crises measured", "Average fall", "Crises not yet recovered", "Avg trading days to recover",
                    "Avg 1-year return after bottom", "Today vs latest crisis peak",
                    "Resilience score (avg fall - 5% per unrecovered crisis)", "Resilience rank (1 = best)"],
           col=sc, height=75)
    for i, ctry in enumerate(countries):
        r = a0 + 1 + i
        ws.cell(r, sc, f'=COUNTIF({rng("A")},$A{r})')
        ws.cell(r, sc + 1, f'=IFERROR(AVERAGEIFS({rng("F")},{rng("A")},$A{r}),"")')
        ws.cell(r, sc + 2, f'=COUNTIFS({rng("A")},$A{r},{rng("J")},0)')
        ws.cell(r, sc + 3, f'=IFERROR(AVERAGEIFS({rng("I")},{rng("A")},$A{r},{rng("J")},1),"")')
        ws.cell(r, sc + 4, f'=IFERROR(AVERAGEIFS({rng("K")},{rng("A")},$A{r}),"")')
        ws.cell(r, sc + 5, f'=IFERROR(AVERAGEIFS({rng("L")},{rng("A")},$A{r},{rng("C")},{L(1 + nc)}${a0}),"")')
        avgf, notrec = f"{L(sc + 1)}{r}", f"{L(sc + 2)}{r}"
        ws.cell(r, sc + 6, f'=IF({avgf}="","",{avgf}-0.05*{notrec})')
        scr = f"${L(sc + 6)}${a0 + 1}:${L(sc + 6)}${aend}"
        ws.cell(r, sc + 7, f'=IF({L(sc + 6)}{r}="","",RANK({L(sc + 6)}{r},{scr},0))')
    style_range(ws, f"{L(sc)}{a0 + 1}:{L(sc + 7)}{aend}")
    for k in (1, 4, 5, 6):
        style_range(ws, f"{L(sc + k)}{a0 + 1}:{L(sc + k)}{aend}", PCT)
    style_range(ws, f"{L(sc + 3)}{a0 + 1}:{L(sc + 3)}{aend}", NUM0)
    ws.conditional_formatting.add(f"{L(sc + 5)}{a0 + 1}:{L(sc + 5)}{aend}", scale3())

    b0 = aend + 4
    bend = matrix(b0, "B. Trading days from bottom to full recovery ('Not yet' = never regained the peak)", "I", NUM0, True)
    ws.conditional_formatting.add(f"B{b0 + 1}:{L(1 + nc)}{bend}", CellIsRule(operator="equal", formula=['"Not yet"'], fill=RED_FILL))
    c0 = bend + 4
    cend = matrix(c0, "C. Return in the year after the bottom (strength of the rebound)", "K", PCT)
    ws.conditional_formatting.add(f"B{c0 + 1}:{L(1 + nc)}{cend}", scale3())
    ws.cell(cend + 2, 1, "Rank method: average fall, with each crisis the market never recovered from counted as an "
                         "extra 5% fall. Rank 1 = most resilient market.").font = F_NOTE
    ws.freeze_panes = "B5"
    widths(ws, {"A": 22})
    for i in range(2, sc + 9):
        ws.column_dimensions[L(i)].width = 12


# ---------------------------------------------------------------- strategy shifts
SHIFTS = [
    ("Q1 2022", "Inflation, Fed hikes, Ukraine war", "Foreign investors sold a record net $13.5 billion of Indian equities in the "
     "first quarter of 2022.", "FPIs exit emerging markets when the Fed tightens and oil rises; DIIs absorbed the selling.",
     "https://www.business-standard.com/article/markets/street-signs-record-fpi-sell-off-nifty-nears-resistance-zone-and-more-122062600651_1.html"),
    ("Oct 2024", "'Buy China, sell India'", "About $10 billion of foreign money left Indian equities in October 2024 as China "
     "announced stimulus; strategists such as Jefferies' Chris Wood raised China at India's expense.",
     "Rich valuations plus a cheaper alternative market trigger tactical rotation away from India.",
     "https://invezz.com/de/news/2024/10/21/rekord-abfluss-von-10-milliarden-us-dollar-aus-auslandischen-investmentfonds-erschuttert-den-indischen-aktienmarkt-im-oktober-ist-china-schuld/"),
    ("Mar-Apr 2025", "Jefferies GREED & fear", "Added 2 points to China and cut India and Korea by 1 point each (March); "
     "in April advised cutting US stocks in favour of Europe, China and India.",
     "Global allocators move country weights in small steps; India regained favour as the US tariff shock hit.",
     "https://www.businesstoday.in/amp/markets/story/greed-fear-jefferies-cuts-india-weight-ups-chinas-says-this-on-fed-rate-cut-467101-2025-03-07"),
    ("Mar 2025", "Domestic investors overtake foreigners", "DIIs held 17.62% of NSE-listed companies vs 17.22% for FPIs, the "
     "first time since tracking began in 2009; by Mar-2026 DIIs held a record 20.9% of the Nifty 500.",
     "SIP and insurance money now cushions FII selling: falls are shallower than FII flows alone would suggest.",
     "https://www.business-standard.com/markets/news/diis-surpass-fpis-in-ownership-of-nse-listed-firms-in-march-2025-125050200426_1.html"),
    ("Feb 2026", "BofA Fund Manager Survey", "'Long gold' was the most crowded trade.",
     "Investors were already hedging before the Iran war.",
     "https://investinglive.com/news/ai-bubble-top-tail-risk-long-gold-most-crowded-trade-according-to-bofa-survey-20260217/"),
    ("Mar 2026", "BofA Fund Manager Survey", "Cash rose to about 4.2-4.3%, the largest monthly jump since 2020; equity "
     "overweight cut to a net 37%, commodities kept.", "Defensive turn: cash and commodities up, equities down.",
     "https://www.scmp.com/business/china-business/article/3347009/iran-conflict-drives-fund-managers-slash-risk-and-hoard-cash-bofa-survey-shows"),
    ("Mar 2026", "BofA (Hartnett): winners and losers of a long Iran war", "Winners: oil, US dollar, US tech, global "
     "defence. Losers: oil importers.", "India, as a large oil importer, sits on the losing side of this playbook.",
     "https://www.investing.com/news/stock-market-news/bofas-hartnett-flags-asset-winners-and-losers-from-prolonged-iran-war-4546287"),
    ("Apr 2026", "BofA Fund Manager Survey", "Growth expectations cut by the most in four years; 'long oil' became one of the "
     "most crowded trades; cash 4.3%, still below the 5% contrarian buy signal.",
     "No capitulation yet: fear was high but positioning was not at extreme levels.",
     "https://www.bloomberg.com/news/articles/2026-04-14/investors-slash-growth-views-by-most-in-four-years-bofa-says"),
    ("Apr 2026", "Indian brokerages cut Nifty targets", "Average 12-month Nifty target cut 3.8% (29,899 to 28,748) two months "
     "into the war; one brokerage cut its target to 25,900 from 29,300.",
     "Forecasters cut targets after the fall, not before it (see Forecast_Studies).",
     "https://www.business-standard.com/amp/markets/news/iran-war-impact-nifty-target-cut-2026-brokerages-flag-oil-inflation-risks-india-126042801604_1.html"),
    ("May 2026", "BofA Fund Manager Survey", "Biggest monthly jump in equity allocation since 2001; cyclicals over defensives "
     "at the highest since Jan-2018; 73% named 'long semiconductors' the most crowded trade.",
     "After the ceasefire, managers swung quickly from fear to greed; crowding moved to AI/semiconductors.",
     "https://www.axios.com/2026/05/20/fund-managers-stocks-bofa"),
    ("2026 YTD", "Record FPI selling in India", "Net FPI equity outflows of roughly Rs 2.7 lakh crore by 1-Oct-2026, above any "
     "previous full year; March alone about Rs 1.17 lakh crore (aggregator figures based on NSDL; please verify).",
     "Foreign money left India for oil-exporter safety, the dollar and East Asian AI markets.",
     "https://www.kotakneo.com/news/market-news/fpi-outflow-2026-nears-30-billion-record-foreign-selling/"),
]


def build_strategy(wb, crises):
    ws = wb.create_sheet("Strategy_Shifts")
    title(ws, "Strategy Shifts — where professional money moved in each crisis",
          "Part A is measured from your watchlist (formulas from Crisis_Sectors). Part B lists documented strategy "
          "changes of fund managers and strategists, with sources.")
    ws["A4"] = "A. Measured sector rotation in each crisis (watchlist industries)"
    ws["A4"].font = F_BOLD
    header(ws, 5, ["Crisis", "Held up best in the fall", "Hit hardest in the fall", "Recovery leader (1 year after bottom)",
                   "Recovery laggard"], height=30)
    R = CRISIS_SECTOR_ROWS
    for j, name in enumerate(crises):
        r = 6 + j
        col = L(2 + j)
        ws.cell(r, 1, name).font = F_BOLD
        for c, key in ((2, "best"), (3, "worst"), (4, "lead"), (5, "lag")):
            ws.cell(r, c, f"=Crisis_Sectors!{col}{R[key]}")
    aend = 5 + len(crises)
    style_range(ws, f"B6:E{aend}")
    for r in range(6, aend + 1):
        ws.cell(r, 1).border = BOX
    ws.cell(aend + 1, 1, "Only industries with at least 3 watchlist stocks are ranked (counts on Crisis_Sectors).").font = F_NOTE

    b0 = aend + 4
    ws.cell(b0 - 1, 1, "B. Documented strategy shifts of fund managers and strategists").font = F_BOLD
    header(ws, b0, ["When", "Who / what", "What they did", "Lesson for our system", "Source"], height=30)
    for i, row in enumerate(SHIFTS):
        r = b0 + 1 + i
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = F_BASE, BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 5).hyperlink = row[4]
        ws.cell(r, 5).font = Font(name=FONT, size=9, color="0563C1", underline="single")
    widths(ws, {"A": 40, "B": 30, "C": 60, "D": 50, "E": 40})


# ---------------------------------------------------------------- analyst forecasts
STUDIES = [
    ("Can stock market forecasters forecast?", "Alfred Cowles III (1933), Econometrica",
     "Studied about 7,500 stock recommendations by 16 US financial services (1928-1932), plus fire insurance "
     "companies and financial publications. As a group the services did worse than the average stock by about "
     "1.4% a year, and the best records looked like luck rather than skill.",
     "Even professional forecasts need to be checked against outcomes before we trust them.",
     "https://economics.yale.edu/sites/default/files/2022-08/cowles-forecasters33.pdf"),
    ("Expert Political Judgment", "Philip Tetlock (2005), Princeton University Press",
     "Tracked 284 experts making 82,361 forecasts over about 20 years. On average they were only slightly better "
     "than chance; famous, confident experts did worse; events called 'impossible' happened about 15% of the time.",
     "The louder and more certain a TV expert is during a crisis, the less weight we give the forecast.",
     "https://www.journalofaccountancy.com/issues/2006/mar/bewareexpertpredictions.html"),
    ("Guru Grades", "CXO Advisory Group (2005-2012)",
     "Graded 6,582 public US stock market forecasts by 68 experts. Average accuracy was 47.4%, lower than a "
     "coin toss; individual experts ranged from 20% to 72%.",
     "Direction calls of market gurus are roughly 50/50. Track each source's hit rate (Forecast_Tracker).",
     "https://www.cxoadvisory.com/gurus/"),
    ("How well do economists forecast recessions?", "An, Jalles and Loungani (2018), IMF Working Paper 18/39",
     "Studied GDP forecasts for 63 countries (1992-2014), covering 153 recessions. The vast majority were missed; "
     "forecasters underestimated how deep recessions would be until the year was nearly over. Private-sector and "
     "official (IMF) forecasts were equally poor.",
     "Do not wait for economists to confirm a recession: by then the market has usually fallen. Use the Panic Meter.",
     "https://www.imf.org/en/publications/wp/issues/2018/03/05/how-well-do-economists-forecast-recessions-45672"),
    ("Equity analysts: Still too bullish", "McKinsey & Company (2010), Goedhart, Raj and Saxena",
     "Over 25 years, analysts forecast S&P 500 earnings growth of 10-12% a year against actual growth of about 6%. "
     "Forecasts were too high in almost every year except recoveries after recessions, and analysts were slow to "
     "cut estimates when the economy weakened.",
     "Discount consensus EPS growth before using it in targets, especially going into a slowdown.",
     "https://www.mckinsey.com/capabilities/strategy-and-corporate-finance/our-insights/equity-analysts-still-too-bullish"),
    ("SPIVA India Scorecard", "S&P Dow Jones Indices (Year-End 2024)",
     "Over 10 years, 74% of Indian large-cap active funds underperformed their benchmark (S&P India LargeMidCap). "
     "Earlier scorecards show similar majorities (about 54% to 68%).",
     "Most professional money managers do not beat the index over time, so 'expert' views are not an edge in themselves.",
     "https://www.spglobal.com/spdji/en/spiva/article/spiva-india-year-end-2024/"),
    ("BofA Global Fund Manager Survey: cash rule", "Bank of America (monthly survey)",
     "When fund managers' average cash level rises above about 5% (4.5% in earlier versions) BofA treats it as a "
     "contrarian BUY signal; below about 3.5-4% as a SELL signal. Since 2011 the buy signal was followed by "
     "average S&P 500 gains of about 7% over six months.",
     "When professionals are most fearful (high cash), returns afterwards have tended to be good: fear is a signal.",
     "https://www.ii.co.uk/analysis-commentary/professional-investor-pessimism-triggers-buy-signal-ii529538"),
    ("Prospect theory and loss aversion", "Kahneman and Tversky (1979), Econometrica",
     "People feel losses roughly twice as strongly as equal gains. This explains panic selling in crashes and "
     "holding losers too long.",
     "Panic is a predictable human reaction, which is why fear peaks tend to overshoot the real economic damage.",
     "https://doi.org/10.2307/1914185"),
]


def build_studies(wb):
    ws = wb.create_sheet("Forecast_Studies")
    title(ws, "Forecast Studies — what research says about expert predictions, and the lesson for our point system",
          "Summaries of published research. Please read the sources for the full findings.")
    header(ws, 4, ["Study", "Who / when", "Main finding", "Lesson for our system", "Source"], height=30)
    for r, row in enumerate(STUDIES, start=5):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(r, c, v)
            cell.font, cell.border = F_BASE, BOX
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).font = F_BOLD
        ws.cell(r, 5).hyperlink = row[4]
        ws.cell(r, 5).font = Font(name=FONT, size=9, color="0563C1", underline="single")
    widths(ws, {"A": 26, "B": 26, "C": 70, "D": 50, "E": 40})
    r = 6 + len(STUDIES)
    ws.cell(r, 1, "Summary").font = Font(name=FONT, size=11, bold=True, color="1F3864")
    ws.cell(r + 1, 1, "Experts rarely predict crises in time and are usually too optimistic before them and too "
                      "pessimistic at the bottom. So we use their forecasts as a SENTIMENT reading, not as a target: "
                      "extreme gloom is scored as fear (often near a bottom), extreme confidence as greed. "
                      "Log forecasts on the Forecast_Tracker sheet to measure each source's real hit rate.").font = F_BASE
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=5)
    ws.cell(r + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r + 1].height = 45


TRACK_ROWS = 500


def build_tracker(wb, mdc, last_md):
    ws = wb.create_sheet("Forecast_Tracker")
    title(ws, "Forecast Tracker — log analyst / brokerage predictions and see how they turned out (formulas)",
          "Enter the yellow columns. The sheet looks up Nifty on the forecast date and at the end of the horizon, "
          "and the Panic Meter reading when the forecast was made. Row 5 is an EXAMPLE: replace it.")
    heads = ["Forecast date", "Source (analyst / brokerage / media)", "Forecast in words", "Predicted direction (Up/Down)",
             "Predicted Nifty level", "Horizon (calendar days)", "Nifty on forecast date", "Panic score on forecast date",
             "Target date", "Nifty at target date", "Actual Nifty change", "Direction correct? (1/0)",
             "Level error", "Status"]
    header(ws, 4, heads, height=60)
    md = lambda k: f"Market_Daily!${mdc[k]}$2:${mdc[k]}${last_md}"  # noqa: E731
    example = [pd.Timestamp("2025-01-02").to_pydatetime(), "Example Broker (replace)",
               "Nifty to reach 26,000 within a year", "Up", 26000, 365]
    end = 4 + TRACK_ROWS
    for i in range(TRACK_ROWS):
        r = 5 + i
        if i == 0:
            for c, v in enumerate(example, start=1):
                ws.cell(r, c, v)
        for c in range(1, 7):
            ws.cell(r, c).fill, ws.cell(r, c).font = FILL_INPUT, F_INPUT
        ws.cell(r, 7, f'=IF(A{r}="","",IFERROR(INDEX({md("Nifty")},MATCH(A{r},{md("Date")},1)),""))')
        ws.cell(r, 8, f'=IF(A{r}="","",IFERROR(INDEX({md("Panic")},MATCH(A{r},{md("Date")},1)),""))')
        ws.cell(r, 9, f'=IF(OR(A{r}="",F{r}=""),"",A{r}+F{r})')
        ws.cell(r, 10, f'=IF(I{r}="","",IF(I{r}>MAX({md("Date")}),"",IFERROR(INDEX({md("Nifty")},MATCH(I{r},{md("Date")},1)),"")))')
        ws.cell(r, 11, f'=IF(OR(G{r}="",J{r}=""),"",J{r}/G{r}-1)')
        ws.cell(r, 12, f'=IF(OR(K{r}="",D{r}=""),"",IF(OR(AND(D{r}="Up",K{r}>0),AND(D{r}="Down",K{r}<0)),1,0))')
        ws.cell(r, 13, f'=IF(OR(E{r}="",J{r}=""),"",E{r}/J{r}-1)')
        ws.cell(r, 14, f'=IF(A{r}="","",IF(J{r}="","Pending",IF(L{r}=1,"Correct","Wrong")))')
    for col, fmt in (("A", "dd-mmm-yyyy"), ("E", "#,##0"), ("G", "#,##0"), ("H", "0"), ("I", "dd-mmm-yyyy"),
                     ("J", "#,##0"), ("K", PCT), ("M", PCT)):
        for c in ws[f"{col}5:{col}{end}"]:
            c[0].number_format = fmt
    dv = DataValidation(type="list", formula1='"Up,Down"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"D5:D{end}")
    ws.conditional_formatting.add(f"N5:N{end}", CellIsRule(operator="equal", formula=['"Correct"'], fill=GREEN_FILL))
    ws.conditional_formatting.add(f"N5:N{end}", CellIsRule(operator="equal", formula=['"Wrong"'], fill=RED_FILL))
    ws.conditional_formatting.add(f"H5:H{end}", panic_scale())
    ws.freeze_panes = "C5"

    # ---- scorecard by source
    header(ws, 4, ["Source (type names to score)", "Forecasts checked", "Hit rate", "Avg level error",
                   "Hit rate when Panic >= 60", "Hit rate when Panic < 40"], col=16, height=60)
    rng = lambda c: f"${c}$5:${c}${end}"  # noqa: E731
    for i in range(15):
        r = 5 + i
        cell = ws.cell(r, 16, "Example Broker (replace)" if i == 0 else None)
        cell.fill, cell.font, cell.border = FILL_INPUT, F_INPUT, BOX
        ws.cell(r, 17, f'=IF(P{r}="","",COUNTIFS({rng("B")},P{r},{rng("L")},"<>"))')
        ws.cell(r, 18, f'=IF(OR(P{r}="",Q{r}=0),"",AVERAGEIFS({rng("L")},{rng("B")},P{r}))')
        ws.cell(r, 19, f'=IF(OR(P{r}="",Q{r}=0),"",IFERROR(AVERAGEIFS({rng("M")},{rng("B")},P{r}),""))')
        ws.cell(r, 20, f'=IF(P{r}="","",IFERROR(AVERAGEIFS({rng("L")},{rng("B")},P{r},{rng("H")},">=60"),""))')
        ws.cell(r, 21, f'=IF(P{r}="","",IFERROR(AVERAGEIFS({rng("L")},{rng("B")},P{r},{rng("H")},"<40"),""))')
    style_range(ws, "Q5:U19")
    for col in "RSTU":
        style_range(ws, f"{col}5:{col}19", PCT)
    ws.cell(21, 16, "A source whose hit rate falls when panic is high is following the crowd's emotion.").font = F_NOTE
    widths(ws, {"A": 12, "B": 24, "C": 36, "D": 11, "E": 11, "F": 10, "G": 11, "H": 10, "I": 12, "J": 11, "K": 11,
                "L": 10, "M": 10, "N": 10, "O": 3, "P": 26})
    for col in "QRSTU":
        ws.column_dimensions[col].width = 12


# ---------------------------------------------------------------- how to use
def build_howto(wb, n_events, n_stocks, first, last):
    ws = wb.create_sheet("How_To_Use", 0)
    title(ws, "Event Impact Analyzer — measuring the market's FEAR and GREED around news")
    lines = [
        ("What this does", None),
        ("", f"Measures how Nifty and your {n_stocks} watchlist stocks reacted to {n_events} events between "
             f"{first} and {last}: elections, budgets, RBI rate decisions, crude oil shocks, bank collapses, "
             "government crises, blackouts and supply shocks, pandemics, global and geopolitical shocks. It also "
             "measures every major recession / bear market from peak to recovery, scores the market's panic level "
             "daily, and checks expert forecasts against outcomes."),
        ("Phases measured", None),
        ("BEFORE", "Run-up in the 20 and 5 trading days before the event (did the market anticipate it?)."),
        ("DURING", "Move on the event day (T0) and the intraday worst point."),
        ("AFTER", "Moves over the next 5, 20 and 60 days; worst fall and best rise within 20 days."),
        ("SETTLING", "Days until the price is back at the pre-event close, and days until India VIX settles."),
        ("R:R", "If you bought 5 days before / on the event day / 5 days after: best gain vs worst dip over 20 days."),
        ("Market mood", "FEAR / GREED / CALM per event, from Nifty's reaction (thresholds on the Settings sheet)."),
        ("Sheets", None),
        ("Settings", "Yellow input cells: fear/greed thresholds, points scale, target weights."),
        ("Event_Calendar", "Every event with the news context and whether it was a surprise."),
        ("Nifty_Impact", "One row per event with all phase measurements for Nifty (Sensex before Sep-2007)."),
        ("Category_Summary", "Averages per event type, positive vs negative news, fear vs greed and who made money. "
                             "Includes Direction and Risk POINTS."),
        ("Stock_Scorecard", "Each watchlist stock: behaviour in fear and greed, Event Beta, resilience and impact ranks."),
        ("Top_Lists", "Top 25 most resilient stocks and top 25 most impacted stocks."),
        ("Sector_Impact", "Industry x event-type heat map."),
        ("Crude_Ranges", "At what size of crude move (and what crude price level) Nifty and industries react."),
        ("FII_DII", "Your daily FII/DII data linked to Nifty moves; paste more history to extend it."),
        ("Target_Adjuster", "Pick a stock and an upcoming event: get adjusted short/long-term targets and stop-loss."),
        ("Panic_Meter", "Daily 0-100 panic score (VIX, drawdown, speed of fall, breadth, rupee) and what Nifty did "
                        "next at each panic level, for days and for events."),
        ("Crisis_Periods", "Every recession / bear market since 1997 including the ongoing 2026 Iran-war crisis: depth, "
                           "duration, recovery, panic peak, buy-early vs wait results, and WHERE ARE WE NOW."),
        ("Crisis_Scorecard, Crisis_Sectors", "How each stock and industry behaved in crises; defenders and "
                                             "recovery leaders are also on Top_Lists."),
        ("Global_Recovery", "11 world markets plus India: fall, recovery time and rebound in every crisis; which never recovered."),
        ("Strategy_Shifts", "Sector rotation measured in each crisis, and documented fund-manager strategy changes."),
        ("Forecast_Studies", "Famous research on how accurate expert and analyst predictions are."),
        ("Forecast_Tracker", "Log analyst predictions: the sheet checks them against what Nifty actually did."),
        ("Stock_Event_Data, Crisis_Stock_Data, Global_Crisis_Data, Market_Daily", "Raw measurements that the formulas read (do not edit)."),
        ("Colour code", None),
        ("Yellow fill, blue text", "Inputs you can change."),
        ("Black text", "Formulas or measured data; do not type over them."),
        ("Green / red shading", "Higher / lower values; FEAR rows red, GREED rows green."),
        ("Data and limits", None),
        ("Prices", "Yahoo Finance daily prices (split-adjusted), downloaded with download_data.py. Yahoo's Nifty "
                   "history starts Sep-2007, so the 1997-2007 events use Sensex (from Jul-1997). One-day data spikes are removed, "
                   "and stock windows containing an unadjusted split/demerger jump are skipped."),
        ("Corporate actions", "Prices are adjusted for splits, bonuses and dividends; buybacks need no adjustment. "
                              "Demergers cannot be adjusted cleanly, so any event or crisis window containing one is skipped."),
        ("Event dates", "Compiled for this analysis from public knowledge: please verify, especially recent and pre-2004 events. "
                        "If an event fell on a holiday or weekend, T0 is the next trading day."),
        ("FII/DII", "Only your Jun-Oct 2026 tracker data is available here; long history must be downloaded from NSE/NSDL."),
        ("Caution", "History shows tendencies, not certainties. Small categories (few events) carry low confidence."),
        ("Updating", None),
        ("", "Add new events to events.csv, then run: python download_data.py --refresh, python analyze.py, "
             "python build_workbook.py."),
    ]
    r = 3
    for a, b in lines:
        if b is None:
            r += 1
            ws.cell(r, 1, a).font = Font(name=FONT, size=11, bold=True, color="1F3864")
        else:
            ws.cell(r, 1, a).font = F_BOLD
            c = ws.cell(r, 2, b)
            c.font = F_BASE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    widths(ws, {"A": 30, "B": 120})


# ---------------------------------------------------------------- main
def main():
    ni = pd.read_csv(DATA / "nifty_impact.csv")
    se = pd.read_csv(DATA / "stock_events.csv")
    md = pd.read_csv(DATA / "market_daily.csv")
    md.columns = [c.replace("/", "_") for c in md.columns]
    watch = pd.read_csv(HERE / "watchlist.csv")

    meta = watch.set_index("NSE_Code")
    se["Industry"] = se["Stock"].map(meta["Industry"])
    se["Cap"] = se["Stock"].map(meta["Cap"])
    ev = ni.set_index("Event_ID")
    se["Category"] = se["Event_ID"].map(ev["Category"])
    se["Nifty_Day0"] = se["Event_ID"].map(ev["Day0"])
    se["Nifty_Total20"] = se["Event_ID"].map(ev["Total20"])

    order = ["Central Election", "Exit Poll", "State Election", "Government Crisis", "Union Budget",
             "Interim Budget", "RBI Rate Hike", "RBI Rate Cut", "RBI Pause", "Crude Spike", "Crude Crash",
             "Bank Collapse", "Pandemic", "Industrial/Supply Shock", "Global Shock", "Geopolitical",
             "Domestic Shock", "Policy Shock", "US Policy"]
    categories = [c for c in order if c in set(ni["Category"])] + \
                 sorted(set(ni["Category"]) - set(order))
    industries = sorted(watch["Industry"].dropna().unique())

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_settings(wb)
    mdc, last_md = build_market_daily(wb, md)
    build_calendar(wb, ni)
    last_ni = build_nifty_impact(wb, ni)
    cat_list = build_category_summary(wb, categories, last_ni)
    last_se = build_stock_events(wb, se, last_ni)
    end_sc = build_scorecard(wb, watch, last_se)
    build_top_lists(wb, end_sc)
    build_sector(wb, industries, categories, last_se)
    build_crude(wb, mdc, last_md)
    build_fii(wb, mdc, last_md)
    build_panic(wb, mdc, last_md, last_ni)
    cp = pd.read_csv(DATA / "crisis_periods.csv")
    cs = pd.read_csv(DATA / "crisis_stocks.csv")
    send = build_crisis(wb, cp, cs, watch, mdc, last_md)
    add_crisis_top_lists(wb, send)
    gc = pd.read_csv(DATA / "global_crises.csv")
    build_global(wb, gc, list(cp["Crisis"]))
    build_strategy(wb, list(cp["Crisis"]))
    build_studies(wb)
    build_tracker(wb, mdc, last_md)

    px = pd.read_csv(DATA / "prices" / "RELIANCE.csv")
    cmp_ = round(float(px["Close"].iloc[-1]), 2)
    build_adjuster(wb, last_se, end_sc, cat_list,
                   {"stock": "RELIANCE", "category": "RBI Rate Hike", "cmp": cmp_,
                    "st": round(cmp_ * 1.08, 2), "lt": round(cmp_ * 1.25, 2)})
    first = pd.Timestamp(ni["Trading_Day"].min()).strftime("%b-%Y")
    last = pd.Timestamp(ni["Trading_Day"].max()).strftime("%b-%Y")
    build_howto(wb, len(ni), len(watch), first, last)

    order_sheets = ["How_To_Use", "Settings", "Target_Adjuster", "Panic_Meter", "Category_Summary",
                    "Crisis_Periods", "Top_Lists", "Stock_Scorecard", "Crisis_Scorecard", "Sector_Impact",
                    "Crisis_Sectors", "Global_Recovery", "Strategy_Shifts", "Crude_Ranges", "FII_DII", "Forecast_Studies",
                    "Forecast_Tracker",
                    "Event_Calendar", "Nifty_Impact", "Stock_Event_Data", "Crisis_Stock_Data", "Global_Crisis_Data", "Market_Daily"]
    wb._sheets = [wb[n] for n in order_sheets]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.zoomScale = 90
    wb.save(OUT)
    print(f"Saved {OUT.name}: {len(ni)} events, {len(se)} stock-event rows")


if __name__ == "__main__":
    main()
